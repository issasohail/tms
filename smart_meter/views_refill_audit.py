"""Filtered audit and exports for prepaid top-up/refund money commands."""

from datetime import date, datetime, time, timedelta
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth.decorators import login_required, permission_required
from django.core.paginator import Paginator
from django.http import Http404, HttpResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.access import restrict_queryset_to_properties
from properties.models import Property, Unit
from smart_meter.models import Meter, MeterCommand, MeterPrepaidRecharge
from smart_meter.utils.tenants import active_tenant_info_for_units


PERIODS = (
    ("this_month", "This month"),
    ("this_week", "This week"),
    ("last_week", "Last week"),
    ("last_month", "Last month"),
    ("custom", "Custom dates"),
    ("all", "All dates"),
)
DEFINITIVE_FAILURE_COMMANDS = frozenset({"failed", "cancelled", "expired", "timeout", "error"})
ACTIVE_COMMANDS = frozenset({"new", "pending", "waiting_online", "claimed", "sent", "acknowledged", "retry"})
EXPORT_HEADERS = (
    "S/N", "Date / time", "Meter #", "Property", "Unit", "Current tenant",
    "Operation", "Amount", "Before", "After", "Transaction status",
    "Command status", "Acknowledged", "Verified", "Order serial", "Name", "Reason / detail",
)


def _date_range(request):
    today = timezone.localdate()
    period = request.GET.get("period") or "this_month"
    if period not in {value for value, _label in PERIODS}:
        period = "this_month"
    if period == "all":
        return period, None, None
    if period == "this_week":
        start = today - timedelta(days=today.weekday())
        return period, start, start + timedelta(days=6)
    if period == "last_week":
        end = today - timedelta(days=today.weekday() + 1)
        return period, end - timedelta(days=6), end
    if period == "last_month":
        end = today.replace(day=1) - timedelta(days=1)
        return period, end.replace(day=1), end
    if period == "custom":
        try:
            start = date.fromisoformat(request.GET.get("from_date") or "")
        except ValueError:
            start = today.replace(day=1)
        try:
            end = date.fromisoformat(request.GET.get("to_date") or "")
        except ValueError:
            end = today
        if start > end:
            start, end = end, start
        return period, start, end
    return "this_month", today.replace(day=1), today


def _aware_start(day):
    return timezone.make_aware(datetime.combine(day, time.min), timezone.get_current_timezone())


def _filtered_transactions(request):
    meters = restrict_queryset_to_properties(
        Meter.objects.select_related("unit", "unit__property"),
        request.user,
        "unit__property",
    )
    period, from_date, to_date = _date_range(request)
    property_id = (request.GET.get("property") or "").strip()
    unit_id = (request.GET.get("unit") or "").strip()
    meter_id = (request.GET.get("meter") or "").strip()

    transactions = MeterPrepaidRecharge.objects.filter(
        pilot__meter_id__in=meters.values("pk")
    ).select_related(
        "pilot__meter", "pilot__meter__unit", "pilot__meter__unit__property", "created_by"
    ).order_by("-created_at", "-pk")
    if from_date:
        transactions = transactions.filter(created_at__gte=_aware_start(from_date))
    if to_date:
        transactions = transactions.filter(created_at__lt=_aware_start(to_date + timedelta(days=1)))
    if property_id.isdigit():
        transactions = transactions.filter(pilot__meter__unit__property_id=int(property_id))
    if unit_id.isdigit():
        transactions = transactions.filter(pilot__meter__unit_id=int(unit_id))
    if meter_id.isdigit():
        transactions = transactions.filter(pilot__meter_id=int(meter_id))

    properties = Property.objects.filter(
        pk__in=meters.values("unit__property_id")
    ).order_by("property_name")
    units = Unit.objects.filter(
        pk__in=meters.values("unit_id")
    ).select_related("property").order_by("property__property_name", "unit_number")
    meter_choices = meters.order_by("unit__property__property_name", "unit__unit_number", "meter_number")
    if property_id.isdigit():
        units = units.filter(property_id=int(property_id))
        meter_choices = meter_choices.filter(unit__property_id=int(property_id))
    if unit_id.isdigit():
        meter_choices = meter_choices.filter(unit_id=int(unit_id))

    return {
        "queryset": transactions,
        "period": period,
        "from_date": from_date,
        "to_date": to_date,
        "property_id": property_id,
        "unit_id": unit_id,
        "meter_id": meter_id,
        "properties": properties,
        "units": units,
        "meters": meter_choices,
    }


def _rows(transactions, start_number=1):
    transactions = list(transactions)
    order_keys = [f"prepaid-order:{item.transaction_id}" for item in transactions]
    commands = {
        command.idempotency_key: command
        for command in MeterCommand.objects.filter(idempotency_key__in=order_keys)
    }
    tenant_info = active_tenant_info_for_units(
        item.pilot.meter.unit_id for item in transactions
    )
    rows = []
    for number, item in enumerate(transactions, start_number):
        meter = item.pilot.meter
        command = commands.get(f"prepaid-order:{item.transaction_id}")
        is_refund = bool(command and command.command_type == "prepaid_refund")
        command_status = command.status if command else "missing"
        info = tenant_info.get(meter.unit_id, {})
        user_name = "System / payment"
        if item.created_by:
            user_name = item.created_by.get_full_name().strip() or item.created_by.get_username()
        can_resend = bool(
            item.status == "failed"
            and command
            and command.status in DEFINITIVE_FAILURE_COMMANDS
        )
        detail_parts = []
        for detail in (
            command.reason if command else "",
            command.error if command else "",
            item.reconciliation_note,
        ):
            detail = (detail or "").strip()
            if detail and detail not in detail_parts:
                detail_parts.append(detail)
        rows.append({
            "number": number,
            "transaction": item,
            "command": command,
            "meter": meter,
            "property_name": getattr(getattr(meter.unit, "property", None), "property_name", "Unassigned"),
            "unit_name": meter.display_location_name,
            "tenant_name": info.get("name", "Vacant"),
            "operation": "Refund" if is_refund else "Top up",
            "operation_key": "refund" if is_refund else "recharge",
            "command_status": command_status,
            "command_status_label": command.get_status_display() if command else "Command missing",
            "acknowledged": bool(command and (command.acknowledged_at or command.raw_ack_hex)),
            "verified": bool(command and command.status == "verified" and item.status == "verified"),
            "user_name": user_name,
            "reason": " | ".join(detail_parts),
            "can_resend": can_resend,
            "is_active": command_status in ACTIVE_COMMANDS or item.status in {"pending", "uncertain"},
            "status_url": (
                reverse("smart_meter:prepaid_money_command_status", args=[command.pk])
                if command else ""
            ),
        })
    return rows


def _filter_summary(data):
    period_label = dict(PERIODS).get(data["period"], data["period"])
    pieces = [f"Period: {period_label}"]
    if data["from_date"] or data["to_date"]:
        pieces.append(f"Dates: {data['from_date'] or 'first'} to {data['to_date'] or 'latest'}")
    selected_property = next((item for item in data["properties"] if str(item.pk) == data["property_id"]), None)
    selected_unit = next((item for item in data["units"] if str(item.pk) == data["unit_id"]), None)
    selected_meter = next((item for item in data["meters"] if str(item.pk) == data["meter_id"]), None)
    if selected_property:
        pieces.append(f"Property: {selected_property.property_name}")
    if selected_unit:
        pieces.append(f"Unit: {selected_unit.unit_number}")
    if selected_meter:
        pieces.append(f"Meter: {selected_meter.meter_number}")
    return " | ".join(pieces)


def _export_values(row):
    item, command = row["transaction"], row["command"]
    return (
        row["number"], timezone.localtime(item.created_at).strftime("%Y-%m-%d %H:%M:%S"),
        row["meter"].meter_number, row["property_name"], row["unit_name"], row["tenant_name"],
        row["operation"], item.amount, item.before_balance, item.after_balance,
        item.get_status_display(), row["command_status_label"],
        "Yes" if row["acknowledged"] else "No", "Yes" if row["verified"] else "No",
        item.transaction_id, row["user_name"], row["reason"],
    )


@login_required
@permission_required("smart_meter.view_meter", raise_exception=True)
def refill_audit(request):
    data = _filtered_transactions(request)
    paginator = Paginator(data["queryset"], 50)
    page = paginator.get_page(request.GET.get("page"))
    rows = _rows(page.object_list, page.start_index() if paginator.count else 1)
    query = request.GET.copy()
    query.pop("page", None)
    return render(request, "smart_meter/refill_audit.html", {
        **data,
        "rows": rows,
        "page_obj": page,
        "periods": PERIODS,
        "filter_query": query.urlencode(),
        "filter_summary": _filter_summary(data),
    })


@login_required
@permission_required("smart_meter.view_meter", raise_exception=True)
def refill_audit_export(request, format):
    if format not in {"xlsx", "pdf", "jpg"}:
        raise Http404
    data = _filtered_transactions(request)
    rows = _rows(data["queryset"], 1)
    values = [_export_values(row) for row in rows]
    generated = timezone.localtime().strftime("Generated %Y-%m-%d %H:%M:%S %Z")
    subtitle = _filter_summary(data)
    filename = f"refill-audit-{timezone.localdate().isoformat()}"

    if format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Refill Audit"
        sheet.append(["Refill Audit"])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(EXPORT_HEADERS))
        sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="17324D")
        sheet.append([subtitle])
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(EXPORT_HEADERS))
        sheet.append([generated])
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(EXPORT_HEADERS))
        sheet.append(list(EXPORT_HEADERS))
        for cell in sheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0D6EFD")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in values:
            sheet.append(list(row))
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:Q{max(4, sheet.max_row)}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.oddFooter.left.text = generated
        sheet.oddFooter.right.text = "Page &P of &N"
        for index, column in enumerate(sheet.columns, 1):
            width = min(34, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width
        output = BytesIO()
        book.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    elif format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfgen import canvas as reportlab_canvas
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        class NumberedCanvas(reportlab_canvas.Canvas):
            def __init__(self, *args, **kwargs):
                reportlab_canvas.Canvas.__init__(self, *args, **kwargs)
                self._saved_page_states = []

            def showPage(self):
                self._saved_page_states.append(dict(self.__dict__))
                self._startPage()

            def save(self):
                page_count = len(self._saved_page_states)
                for page_state in self._saved_page_states:
                    self.__dict__.update(page_state)
                    self.setFont("Helvetica", 7)
                    self.drawRightString(
                        landscape(A3)[0] - 24,
                        14,
                        f"Page {self._pageNumber} of {page_count}",
                    )
                    reportlab_canvas.Canvas.showPage(self)
                reportlab_canvas.Canvas.save(self)

        output = BytesIO()
        styles = getSampleStyleSheet()
        document = SimpleDocTemplate(
            output, pagesize=landscape(A3), leftMargin=24, rightMargin=24,
            topMargin=30, bottomMargin=30,
        )
        table_data = [list(EXPORT_HEADERS)] + [[str(value or "") for value in row] for row in values]
        table = Table(table_data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D6EFD")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 6),
            ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))

        def footer(canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 7)
            canvas.drawString(24, 14, generated)
            canvas.restoreState()

        story = [Paragraph("Refill Audit", styles["Title"]), Paragraph(subtitle, styles["Normal"]), Spacer(1, 12), table]
        document.build(
            story,
            onFirstPage=footer,
            onLaterPages=footer,
            canvasmaker=NumberedCanvas,
        )
        response = HttpResponse(output.getvalue(), content_type="application/pdf")
    else:
        from PIL import Image, ImageDraw, ImageFont

        per_page = 28
        pages = [values[index:index + per_page] for index in range(0, len(values), per_page)] or [[]]
        rendered = []
        for page_number, page_rows in enumerate(pages, 1):
            image = Image.new("RGB", (2600, 1700), "white")
            draw = ImageDraw.Draw(image)
            font = ImageFont.load_default()
            draw.rectangle((0, 0, 2600, 105), fill="#17324D")
            draw.text((35, 22), "REFILL AUDIT", fill="white", font=font)
            draw.text((35, 52), subtitle, fill="#DCEAF7", font=font)
            widths = [65, 175, 125, 120, 115, 145, 85, 85, 85, 85, 115, 115, 85, 70, 175, 120, 350]
            x_positions, x = [], 20
            for width in widths:
                x_positions.append(x)
                x += width
            y = 125
            draw.rectangle((15, y, 2585, y + 42), fill="#0D6EFD")
            for x, heading in zip(x_positions, EXPORT_HEADERS):
                draw.text((x, y + 13), heading, fill="white", font=font)
            y += 42
            for row_number, row in enumerate(page_rows):
                if row_number % 2:
                    draw.rectangle((15, y, 2585, y + 47), fill="#F1F5F9")
                for x, width, value in zip(x_positions, widths, row):
                    draw.text((x, y + 14), str(value or "")[:max(6, width // 7)], fill="#17202A", font=font)
                y += 47
            draw.text((20, 1670), generated, fill="#475569", font=font)
            footer = f"Page {page_number} of {len(pages)}"
            draw.text((2520 - len(footer) * 7, 1670), footer, fill="#475569", font=font)
            page_output = BytesIO()
            image.save(page_output, format="JPEG", quality=92)
            rendered.append(page_output.getvalue())
        if len(rendered) == 1:
            response = HttpResponse(rendered[0], content_type="image/jpeg")
        else:
            archive = BytesIO()
            with ZipFile(archive, "w", ZIP_DEFLATED) as zipped:
                for page_number, content in enumerate(rendered, 1):
                    zipped.writestr(f"{filename}-page-{page_number:03d}.jpg", content)
            response = HttpResponse(archive.getvalue(), content_type="application/zip")
            format = "zip"
    response["Content-Disposition"] = f'attachment; filename="{filename}.{format}"'
    return response
