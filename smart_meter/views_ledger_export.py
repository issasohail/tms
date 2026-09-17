"""Download the filtered meter ledger in printable and spreadsheet formats."""

from io import BytesIO
from decimal import Decimal

from django.contrib.auth.decorators import login_required, permission_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from accounts.access import restrict_queryset_to_properties
from smart_meter.ledger_display import build_ledger_data
from smart_meter.models import Meter


def _text(value):
    return "" if value is None else str(value)


def _export_sections(data):
    movements = [["S/N", "Reading", "Balance", "Energy", "Balance change", "Usage change", "Rate snapshot", "Estimated charge"]]
    for reading in data["readings"]:
        energy = reading.forward_active_energy_kwh if reading.forward_active_energy_kwh is not None else reading.total_energy
        movements.append([f"{reading.day_number}.{reading.sub_number}", timezone.localtime(reading.ts).strftime("%Y-%m-%d %H:%M:%S"), reading.balance, energy, reading.balance_change, reading.energy_change, reading.unit_rate, reading.estimated_charge])

    observations = [["S/N", "Date", "Balance", "Energy", "Source IP", "Source port", "Trust"]]
    for frame in data["raw_balance_frames"]:
        observations.append([f"{frame.day_number}.{frame.sub_number}", timezone.localtime(frame.received_at).strftime("%Y-%m-%d %H:%M:%S"), frame.reported_balance, frame.reported_energy, frame.source_ip, frame.source_port, frame.get_trust_classification_display()])

    money = [["S/N", "Date", "Operation", "Amount", "Before", "After", "Status", "Order"]]
    for item in data["transactions"]:
        money.append([item.serial_number, timezone.localtime(item.created_at).strftime("%Y-%m-%d %H:%M:%S"), item.operation_label, item.amount, item.before_balance, item.after_balance, item.get_status_display(), item.transaction_id])
    return [("Historical movements", movements), ("Raw observations", observations), ("Money transactions", money)]


@login_required
@permission_required("smart_meter.view_meter", raise_exception=True)
def meter_ledger_export(request, meter_id, format):
    if format not in ("pdf", "xlsx", "jpg"):
        from django.http import Http404
        raise Http404
    meters = restrict_queryset_to_properties(Meter.objects.select_related("unit", "unit__property"), request.user, "unit__property")
    meter = get_object_or_404(meters, pk=meter_id)
    data = build_ledger_data(request, meter, export=True)
    sections = _export_sections(data)
    filename = f"meter-ledger-{meter.meter_number}"

    if format == "xlsx":
        book = Workbook()
        book.remove(book.active)
        for title, rows in sections:
            sheet = book.create_sheet(title)
            for row in rows:
                sheet.append([cell if isinstance(cell, (int, float, Decimal)) else _text(cell) for cell in row])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="263238")
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(30, max(12, max(len(_text(cell.value)) for cell in column) + 2))
        output = BytesIO()
        book.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    elif format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet

        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=landscape(A3), leftMargin=30, rightMargin=30)
        styles = getSampleStyleSheet()
        story = [Paragraph(f"Meter Balance Ledger: {meter.meter_number}", styles["Title"]),
                 Paragraph(f"Period: {data['from_date'] or 'first record'} to {data['to_date'] or 'latest record'}", styles["Normal"]), Spacer(1, 12)]
        for title, rows in sections:
            table = Table([[_text(cell) for cell in row] for row in rows], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#263238")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5f8")]),
            ]))
            story.extend([Paragraph(title, styles["Heading2"]), table, Spacer(1, 18)])
        document.build(story)
        response = HttpResponse(output.getvalue(), content_type="application/pdf")
    else:
        from PIL import Image, ImageDraw, ImageFont
        from zipfile import ZipFile, ZIP_DEFLATED

        def draw_jpg(part):
            row_count = sum(len(rows) + 2 for _, rows in part)
            image = Image.new("RGB", (2300, row_count * 19 + 110), "white")
            draw = ImageDraw.Draw(image)
            font = ImageFont.load_default()
            y = 16
            draw.text((12, y), f"Meter Balance Ledger: {meter.meter_number}", fill="#18212b", font=font)
            y += 20
            draw.text((12, y), f"Period: {data['from_date'] or 'first record'} to {data['to_date'] or 'latest record'}", fill="#18212b", font=font)
            y += 38
            for title, rows in part:
                draw.text((12, y), title, fill="#18212b", font=font)
                y += 20
                width = 2276 // len(rows[0])
                for number, row in enumerate(rows):
                    if number == 0:
                        draw.rectangle((8, y, 2292, y + 18), fill="#263238")
                    elif number % 2 == 0:
                        draw.rectangle((8, y, 2292, y + 18), fill="#f2f5f8")
                    for index, value in enumerate(row):
                        draw.text((12 + index * width, y + 3), _text(value)[:43], fill="white" if number == 0 else "#18212b", font=font)
                    y += 19
                y += 19
            output = BytesIO()
            image.crop((0, 0, 2300, y + 8)).save(output, format="JPEG", quality=90)
            return output.getvalue()

        # A JPEG cannot exceed 65,535 pixels in height. Long periods download
        # as a ZIP of numbered JPG pages so no filtered rows are omitted.
        max_rows_per_jpg = 900
        pages, current, count = [], [], 0
        for title, rows in sections:
            header, remaining = rows[0], rows[1:]
            if not remaining:
                current.append((title, [header]))
                count += 3
            while remaining:
                if count >= max_rows_per_jpg - 3:
                    pages.append(current)
                    current, count = [], 0
                take = min(len(remaining), max_rows_per_jpg - count - 3)
                current.append((title, [header] + remaining[:take]))
                remaining = remaining[take:]
                count += take + 3
        if current:
            pages.append(current)
        if len(pages) == 1:
            response = HttpResponse(draw_jpg(pages[0]), content_type="image/jpeg")
        else:
            output = BytesIO()
            with ZipFile(output, "w", ZIP_DEFLATED) as archive:
                for number, part in enumerate(pages, 1):
                    archive.writestr(f"{filename}-{number:03d}.jpg", draw_jpg(part))
            response = HttpResponse(output.getvalue(), content_type="application/zip")
            format = "zip"
    response["Content-Disposition"] = f'attachment; filename="{filename}.{format}"'
    return response
