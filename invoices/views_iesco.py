from __future__ import annotations

import csv
import io
import logging
import re
from datetime import date, datetime, timedelta
from decimal import Decimal

import requests
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from django.views.generic import ListView, TemplateView

from accounts.access import (
    has_all_property_access,
    restrict_queryset_to_properties,
)
from leases.models import Lease
from properties.models import Property, Unit

from .forms import IescoMeterAssignmentForm, IescoStandaloneMeterForm
from .iesco_bill_fetch import VpnDetectedError
from .models import IescoBillReading, IescoStandaloneMeter
from .services_iesco import (
    IESCO_EXPORT_FIELDS,
    IescoInvoiceAmountChangeRequired,
    bill_month_start,
    fetch_bill_payload,
    normalize_bill_payload,
    payload_to_csv_row,
    post_reading_to_invoice,
    save_bill_payload,
    validate_reference_no,
)
from .services_iesco_reminders import (
    build_iesco_reminder_message,
    reading_requires_payment,
    reminder_recipient,
    send_iesco_reminder,
)


logger = logging.getLogger(__name__)
PREVIEW_SESSION_KEY = "iesco_bill_preview"
EXPORT_SESSION_KEY = "iesco_bill_last_export_ids"
IMPORT_PROGRESS_SESSION_KEY = "iesco_bill_import_progress_rows"
MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 500
IESCO_BULK_FETCH_WAIT_DAYS = 20
IESCO_BULK_RETRY_DAYS = 3


def _valid_unit_meters(user, *, property_id=None, active_only=False):
    queryset = (
        Unit.objects.select_related("property")
        .filter(electric_meter_num__regex=r"^[0-9]{14}$")
        .exclude(electric_meter_num="00000000000000")
    )
    queryset = restrict_queryset_to_properties(queryset, user, "property")
    if property_id:
        queryset = queryset.filter(property_id=property_id)
    if active_only:
        queryset = queryset.filter(iesco_bill_active=True)
    return queryset.order_by("property__property_name", "unit_number", "id").distinct()


def _visible_standalone_meters(user, *, active_only=False):
    if not has_all_property_access(user):
        return IescoStandaloneMeter.objects.none()
    queryset = IescoStandaloneMeter.objects.all()
    if active_only:
        queryset = queryset.filter(is_active=True)
    return queryset


def _source_description(reference_no, user):
    units = list(_valid_unit_meters(user).filter(electric_meter_num=reference_no)[:2])
    if units:
        unit = units[0]
        return f"{unit.property.property_name} / {unit.unit_number}"
    standalone = _visible_standalone_meters(user).filter(reference_no=reference_no).first()
    return standalone.description if standalone else ""


def _ensure_reference_access(user, reference_no):
    if _valid_unit_meters(user).filter(electric_meter_num=reference_no).exists():
        return
    if _visible_standalone_meters(user).filter(reference_no=reference_no).exists():
        return
    raise Http404("IESCO meter not found")


def _require_change_permission(user):
    if not (user.is_superuser or user.has_perm("invoices.change_iescobillreading")):
        raise PermissionDenied("You do not have permission to fetch or import IESCO bills.")


@require_POST
@login_required
def set_reading_trust(request, pk):
    """Explicitly move a saved IESCO reading through Verified -> Confirmed.

    Re-fetching/importing the same bill resets it to Parsed, so reconciliation
    never silently consumes changed external data.
    """
    _require_change_permission(request.user)
    reading = get_object_or_404(IescoBillReading, pk=pk)
    _ensure_reference_access(request.user, reading.reference_no)
    action = (request.POST.get("action") or "").strip().lower()
    now = timezone.now()
    if action == "verify":
        reading.trust_status = IescoBillReading.TRUST_VERIFIED
        reading.verified_at = now
        reading.verified_by = request.user
        reading.confirmed_at = None
        reading.confirmed_by = None
        reading.save(update_fields=["trust_status", "verified_at", "verified_by", "confirmed_at", "confirmed_by", "updated_at"])
        messages.success(request, "IESCO reading verified. Confirm it separately before reconciliation can use it.")
    elif action == "confirm":
        if reading.trust_status != IescoBillReading.TRUST_VERIFIED or not reading.verified_at:
            messages.error(request, "Verify this IESCO reading before confirming it for reconciliation.")
        else:
            reading.trust_status = IescoBillReading.TRUST_CONFIRMED
            reading.confirmed_at = now
            reading.confirmed_by = request.user
            reading.save(update_fields=["trust_status", "confirmed_at", "confirmed_by", "updated_at"])
            messages.success(request, "IESCO reading confirmed for energy reconciliation.")
    elif action == "reopen":
        reading.trust_status = IescoBillReading.TRUST_PARSED
        reading.verified_at = None
        reading.verified_by = None
        reading.confirmed_at = None
        reading.confirmed_by = None
        reading.save(update_fields=["trust_status", "verified_at", "verified_by", "confirmed_at", "confirmed_by", "updated_at"])
        messages.warning(request, "IESCO reading reopened; reconciliation will not use it until re-verified and confirmed.")
    else:
        messages.error(request, "Unknown IESCO trust action.")
    return redirect("invoices:iesco_bill_reading_detail", reference_no=reading.reference_no)


def _validation_message(exc):
    if isinstance(exc, VpnDetectedError):
        return str(exc)
    if isinstance(exc, (requests.exceptions.SSLError, requests.exceptions.ConnectionError)):
        return (
            "PITC could not be reached from this server. No data was changed. "
            "Use Fetch All Active on the local Pakistani server, download Bills CSV, "
            "then use Upload Bills CSV on this page to update production."
        )
    if hasattr(exc, "messages"):
        return "; ".join(exc.messages)
    return str(exc)


def _fetch_failure_message(exc):
    message = _validation_message(exc)
    if isinstance(exc, VpnDetectedError):
        return message
    if "local Pakistani server" in message:
        return message
    return (
        f"{message} Bill data was not changed. If this is the "
        "hosted production TMS, PITC may reject overseas requests: download the "
        "active reference CSV here, fetch it on the local Pakistan TMS, download "
        "the bills CSV there, and upload that bills CSV back on production."
    )


def _clean_error_message(exc):
    if isinstance(exc, VpnDetectedError):
        return str(exc)
    if isinstance(exc, (requests.exceptions.SSLError, requests.exceptions.ConnectionError)):
        return (
            "Could not connect to PITC / IESCO server. Please check internet connection or retry."
        )
    if isinstance(exc, requests.exceptions.Timeout):
        return "Connection to PITC / IESCO timed out. The server took too long to respond."
    if isinstance(exc, requests.exceptions.HTTPError):
        status_code = exc.response.status_code if getattr(exc, "response", None) is not None else ""
        return f"IESCO server returned an error (HTTP {status_code})." if status_code else "IESCO server returned an HTTP error."
    if hasattr(exc, "messages"):
        return "; ".join(exc.messages)
    msg = str(exc).strip()
    if "without the required bill fields" in msg:
        return "No bill found on PITC for this reference number (or bill page format changed)."
    if "exactly 14 digits" in msg:
        return "Invalid reference number: must be exactly 14 digits."
    return msg


def _month_sort_key(value):
    try:
        return bill_month_start(value)
    except ValidationError:
        return date.min


def _month_options_for_references(references):
    unique = {}
    raw_months = IescoBillReading.objects.filter(
        reference_no__in=references
    ).values_list("bill_month", flat=True).distinct()
    for raw_month in raw_months:
        try:
            month_date = bill_month_start(raw_month)
        except ValidationError:
            continue
        unique[month_date] = {
            "value": month_date.strftime("%Y-%m"),
            "label": month_date.strftime("%B %Y"),
        }
    return [unique[key] for key in sorted(unique, reverse=True)]


def _raw_month_values(references, selected):
    if not re.fullmatch(r"\d{4}-\d{2}", selected or ""):
        return [selected]
    try:
        selected_date = datetime.strptime(selected, "%Y-%m").date().replace(day=1)
    except ValueError:
        return []
    values = []
    raw_months = IescoBillReading.objects.filter(
        reference_no__in=references
    ).values_list("bill_month", flat=True).distinct()
    for raw_month in raw_months:
        try:
            if bill_month_start(raw_month) == selected_date:
                values.append(raw_month)
        except ValidationError:
            continue
    return values



IESCO_PERIOD_CHOICES = (
    ("this_month", "This month"),
    ("last_month", "Last month"),
    ("this_quarter", "This quarter"),
    ("last_quarter", "Last quarter"),
    ("this_year", "This year"),
    ("last_year", "Last year"),
)


def _period_bounds(code, today=None):
    today = today or timezone.localdate()
    first_this_month = today.replace(day=1)
    if code == "this_month":
        start = first_this_month
        end = (
            start.replace(year=start.year + 1, month=1)
            if start.month == 12
            else start.replace(month=start.month + 1)
        ) - timedelta(days=1)
        return start, end
    if code == "last_month":
        end = first_this_month - timedelta(days=1)
        return end.replace(day=1), end
    if code in {"this_quarter", "last_quarter"}:
        quarter_start_month = ((today.month - 1) // 3) * 3 + 1
        start = today.replace(month=quarter_start_month, day=1)
        if code == "last_quarter":
            end = start - timedelta(days=1)
            start_month = ((end.month - 1) // 3) * 3 + 1
            start = end.replace(month=start_month, day=1)
            return start, end
        if quarter_start_month == 10:
            end = start.replace(year=start.year + 1, month=1) - timedelta(days=1)
        else:
            end = start.replace(month=quarter_start_month + 3) - timedelta(days=1)
        return start, end
    if code == "this_year":
        return date(today.year, 1, 1), date(today.year, 12, 31)
    if code == "last_year":
        year = today.year - 1
        return date(year, 1, 1), date(year, 12, 31)
    return None, None


def _period_matches(reading, code):
    start, end = _period_bounds(code)
    if not start or not end:
        return True
    try:
        month = bill_month_start(reading.bill_month)
    except ValidationError:
        return False
    return start <= month <= end


def _latest_saved_month_value(readings):
    months = []
    for reading in readings:
        try:
            months.append(bill_month_start(reading.bill_month))
        except ValidationError:
            continue
    if not months:
        return ""
    return max(months).strftime("%Y-%m")


def _reading_amount_decimal(value):
    cleaned = re.sub(r"[^0-9.\-]", "", str(value or "").replace(",", ""))
    if not cleaned or cleaned in {"-", ".", "-."}:
        return Decimal("0")
    try:
        return Decimal(cleaned)
    except Exception:
        return Decimal("0")


def _display_decimal(value, *, money=False):
    value = value or Decimal("0")
    if money:
        return f"Rs. {value:,.0f}"
    rendered = f"{value:,.2f}"
    return rendered.rstrip("0").rstrip(".")


def _dashboard_totals(rows):
    total_units = Decimal("0")
    grand_total = Decimal("0")
    for row in rows:
        reading = row.get("reading")
        if not reading:
            continue
        imported = reading.import_units
        if imported is None:
            imported = _reading_amount_decimal(reading.units)
        total_units += imported or Decimal("0")
        grand_total += _reading_amount_decimal(reading.grand_total)
    return {
        "units": total_units,
        "units_display": _display_decimal(total_units),
        "grand_total": grand_total,
        "grand_total_display": _display_decimal(grand_total, money=True),
    }


def _visible_sources_for_filters(user, *, property_id="", unit_id="", reference_no=""):
    units = _valid_unit_meters(
        user,
        property_id=property_id if property_id.isdigit() else None,
    )
    if unit_id.isdigit():
        units = units.filter(pk=unit_id)
    if reference_no:
        units = units.filter(electric_meter_num__icontains=reference_no)

    active_units = _valid_unit_meters(
        user,
        property_id=property_id if property_id.isdigit() else None,
        active_only=True,
    )
    if unit_id.isdigit():
        active_units = active_units.filter(pk=unit_id)
    active_unit_references = set(
        active_units.values_list("electric_meter_num", flat=True)
    )

    sources = [
        {
            "reference_no": unit.electric_meter_num,
            "property": unit.property,
            "unit": unit,
            "standalone": None,
            "description": "",
            "is_standalone": False,
            "is_active": unit.electric_meter_num in active_unit_references,
        }
        for unit in units
    ]
    if not property_id and not unit_id:
        standalone = _visible_standalone_meters(user)
        if reference_no:
            standalone = standalone.filter(reference_no__icontains=reference_no)
        standalone = standalone.exclude(
            reference_no__in={row["reference_no"] for row in sources}
        )
        sources.extend(
            {
                "reference_no": meter.reference_no,
                "property": None,
                "unit": None,
                "standalone": meter,
                "description": meter.description,
                "is_standalone": True,
                "is_active": meter.is_active,
            }
            for meter in standalone
        )
    return sources


def _select_dashboard_readings(
    readings,
    *,
    references,
    bill_month="",
    period="",
    period_explicit=False,
    history_mode=False,
):
    ordered = list(readings)
    effective_bill_month = bill_month

    if bill_month:
        raw_values = set(_raw_month_values(references, bill_month))
        ordered = [reading for reading in ordered if reading.bill_month in raw_values]
    elif history_mode:
        if period_explicit and period:
            ordered = [reading for reading in ordered if _period_matches(reading, period)]
    else:
        effective_bill_month = _latest_saved_month_value(ordered)
        if effective_bill_month:
            raw_values = set(_raw_month_values(references, effective_bill_month))
            ordered = [reading for reading in ordered if reading.bill_month in raw_values]

    ordered.sort(
        key=lambda reading: (
            _month_sort_key(reading.bill_month),
            reading.updated_at,
            reading.pk,
        ),
        reverse=True,
    )

    if not history_mode:
        latest = {}
        for reading in ordered:
            latest.setdefault(reading.reference_no, reading)
        ordered = list(latest.values())
    return ordered, effective_bill_month


def _payment_status_matches(reading, payment_status):
    if payment_status == "paid":
        return bool(reading and reading.current_month_paid is True)
    if payment_status == "unpaid":
        return bool(reading and reading_requires_payment(reading))
    if payment_status == "unknown":
        return bool(not reading or reading.current_month_paid is None)
    return True


def _set_preview(request, payloads, errors, *, source):
    deduplicated = {}
    for payload in payloads:
        normalized = normalize_bill_payload(payload)
        display_reading = IescoBillReading(
            units=normalized.get("units"),
            meter_readings=normalized.get("meter_readings"),
            current_bill=normalized.get("current_bill"),
            arrears=normalized.get("arrears"),
            grand_total=normalized.get("grand_total"),
            current_month_paid=normalized.get("current_month_paid"),
        )
        normalized["units_display"] = display_reading.units_display
        normalized["per_unit_rate_display"] = display_reading.per_unit_rate_display
        normalized["current_bill_display"] = display_reading.current_bill_display
        normalized["arrears_display"] = display_reading.arrears_display
        normalized["grand_total_display"] = display_reading.grand_total_display
        normalized["payment_status_display"] = display_reading.payment_status_display
        deduplicated[(normalized["reference_no"], normalized["bill_month"])] = normalized
    request.session[PREVIEW_SESSION_KEY] = {
        "payloads": list(deduplicated.values()),
        "errors": errors,
        "source": source,
        "created_at": timezone.now().isoformat(),
    }
    request.session.modified = True


def _active_fetch_sources(user, property_id=""):
    units = _valid_unit_meters(
        user,
        property_id=property_id if property_id.isdigit() else None,
        active_only=True,
    )
    sources = {
        unit.electric_meter_num: f"{unit.property.property_name} / {unit.unit_number}"
        for unit in units
    }
    if not property_id:
        for meter in _visible_standalone_meters(user, active_only=True):
            sources.setdefault(meter.reference_no, meter.description)
    return sources


def _parsed_iesco_date(value):
    cleaned = re.sub(r"\s+", " ", str(value or "").strip())
    for fmt in ("%d %b %y", "%d-%b-%y", "%d %B %y", "%Y-%m-%d"):
        try:
            return datetime.strptime(cleaned.title(), fmt).date()
        except ValueError:
            continue
    return None


def _bulk_fetch_wait_until(reference_no):
    latest = (
        IescoBillReading.objects.filter(reference_no=reference_no)
        .order_by("-fetched_at", "-updated_at", "-id")
        .first()
    )
    if latest is None:
        return None
    reading_date = _parsed_iesco_date(latest.reading_date)
    fetched_date = (
        timezone.localtime(latest.fetched_at).date()
        if latest.fetched_at
        else timezone.localtime(latest.updated_at).date()
    )
    cycle_date = (
        reading_date + timedelta(days=IESCO_BULK_FETCH_WAIT_DAYS)
        if reading_date
        else fetched_date + timedelta(days=IESCO_BULK_FETCH_WAIT_DAYS)
    )
    if reading_requires_payment(latest) or latest.current_month_paid is None:
        cycle_date = fetched_date + timedelta(days=IESCO_BULK_RETRY_DAYS)
    retry_date = fetched_date + timedelta(days=IESCO_BULK_RETRY_DAYS)
    wait_until = max(cycle_date, retry_date)
    return wait_until if timezone.localdate() < wait_until else None


def _demo_history_readings(reference_no, base_reading=None):
    samples = [
        ("JUL 26", 162, 3400, 0, True, 8, 24),
        ("JUN 26", 148, 3125, 450, False, 9, 25),
        ("MAY 26", 139, 2910, 0, True, 8, 24),
        ("APR 26", 151, 3200, 0, True, 8, 24),
    ]
    base_registers = list(base_reading.meter_readings or []) if base_reading else []
    meter_no = base_registers[0].get("meter_no") if base_registers else "DEMO-METER"
    is_three_phase = bool(base_reading and base_reading.meter_type == "3-P")
    readings = []
    cumulative = Decimal("1000")
    for index, (month, units, current_bill, arrears, paid, reading_day, due_day) in enumerate(samples):
        month_date = bill_month_start(month)
        if is_three_phase:
            import_off_peak = Decimal(units) * Decimal("0.75")
            import_peak = Decimal(units) - import_off_peak
            export_off_peak = Decimal(units) * Decimal("0.18")
            export_peak = Decimal(units) * Decimal("0.04")
            registers = []
            for direction, period, register_units in (
                ("import", "off_peak", import_off_peak),
                ("import", "peak", import_peak),
                ("export", "off_peak", export_off_peak),
                ("export", "peak", export_peak),
            ):
                previous = cumulative
                present = previous + register_units
                registers.append(
                    {
                        "direction": direction,
                        "period": period,
                        "meter_no": meter_no,
                        "previous": f"{previous:.2f}",
                        "present": f"{present:.2f}",
                        "units": f"{register_units:.2f}",
                    }
                )
                cumulative = present
        else:
            previous = cumulative
            present = previous + Decimal(units)
            registers = [
                {
                    "direction": "import",
                    "period": "off_peak",
                    "meter_no": meter_no,
                    "previous": f"{previous:.2f}",
                    "present": f"{present:.2f}",
                    "units": str(units),
                }
            ]
            cumulative = present
        saved_at = timezone.make_aware(datetime(2026, month_date.month, 10, 12, index))
        reading = IescoBillReading(
            reference_no=reference_no,
            bill_month=month,
            meter_type="3-P" if is_three_phase else "S-P",
            meter_readings=registers,
            units=str(units),
            reading_date=f"{reading_day:02d} {month_date.strftime('%b').upper()} 26",
            due_date=f"{due_day:02d} {month_date.strftime('%b').upper()} 26",
            current_bill=f"{current_bill:,}",
            arrears=f"{arrears:,}",
            grand_total=f"{current_bill + arrears:,}",
            current_month_paid=paid,
            amount_paid=f"{current_bill + arrears:,}" if paid else None,
            consumer_name=(base_reading.consumer_name if base_reading else "Demo consumer"),
            tariff_category=(base_reading.tariff_category if base_reading else "Domestic"),
            address=(base_reading.address if base_reading else "Preview data — not saved"),
            fetched_at=saved_at,
            updated_at=saved_at,
        )
        reading.is_demo = True
        readings.append(reading)
    return readings


def _append_preview(request, *, payload=None, error=None):
    batch = request.session.get(PREVIEW_SESSION_KEY) or {}
    payloads = list(batch.get("payloads") or [])
    errors = list(batch.get("errors") or [])
    if payload is not None:
        payloads.append(payload)
    if error is not None:
        errors.append(error)
    _set_preview(request, payloads, errors, source="bulk-fetch")


def _attach_lease_context(rows):
    unit_ids = {row["unit"].pk for row in rows if row.get("unit")}
    leases_by_unit = {}
    if unit_ids:
        leases = (
            Lease.objects.filter(unit_id__in=unit_ids)
            .exclude(status__in=("pending_approval", "rejected"))
            .select_related("tenant")
            .prefetch_related("invoices", "payments__detail", "security_transactions")
            .order_by("unit_id", "-start_date", "-id")
        )
        for lease in leases:
            leases_by_unit.setdefault(lease.unit_id, []).append(lease)

    for row in rows:
        reading = row.get("reading")
        unit = row.get("unit")
        row["lease"] = None
        row["tenant_name"] = ""
        row["lease_balance"] = None
        row["whatsapp_phone"] = ""
        row["invoice"] = (
            reading.posted_invoice_item.invoice
            if reading and reading.posted_invoice_item_id
            else None
        )
        row["invoice_generated"] = bool(
            row["invoice"]
        )
        row["invoice_number"] = (
            row["invoice"].invoice_number
            if row["invoice_generated"]
            else ""
        )
        if reading and not unit and row.get("standalone"):
            row["whatsapp_phone"] = row["standalone"].phone or ""
        if not reading or not unit:
            continue
        row["whatsapp_phone"] = unit.property.owner_phone or ""
        try:
            month_start = bill_month_start(reading.bill_month)
        except ValidationError:
            continue
        month_end = (
            month_start.replace(year=month_start.year + 1, month=1)
            if month_start.month == 12
            else month_start.replace(month=month_start.month + 1)
        ) - timedelta(days=1)
        lease = next(
            (
                candidate
                for candidate in leases_by_unit.get(unit.pk, [])
                if candidate.start_date <= month_end
                and candidate.end_date >= month_start
            ),
            None,
        )
        if lease is None:
            continue
        row["lease"] = lease
        row["tenant_name"] = lease.tenant.get_full_name()
        row["lease_balance"] = lease.get_balance
        if row["invoice"] is None:
            row["invoice"] = next(
                (
                    invoice
                    for invoice in lease.invoices.all()
                    if invoice.issue_date == month_start
                    and invoice.status != "cancelled"
                    and invoice.lifecycle_status not in ("cancelled", "void")
                ),
                None,
            )
            row["invoice_generated"] = bool(row["invoice"])
            row["invoice_number"] = (
                row["invoice"].invoice_number if row["invoice"] else ""
            )
        row["whatsapp_phone"] = (
            lease.tenant.phone or unit.property.owner_phone or ""
        )


def _save_payloads(request, payloads):
    created = 0
    updated = 0
    reading_ids = []
    with transaction.atomic():
        for payload in payloads:
            reference_no = payload["reference_no"]
            try:
                _ensure_reference_access(request.user, reference_no)
            except Http404:
                if not (
                    has_all_property_access(request.user)
                    and (payload.get("description") or "").strip()
                ):
                    raise
            reading, was_created = save_bill_payload(payload)
            reading_ids.append(reading.pk)
            created += int(was_created)
            updated += int(not was_created)
            description = (payload.get("description") or "").strip()
            if description and not Unit.objects.filter(
                electric_meter_num=reference_no
            ).exists():
                IescoStandaloneMeter.objects.update_or_create(
                    reference_no=reference_no,
                    defaults={"description": description, "is_active": True},
                )
    existing_ids = request.session.get(EXPORT_SESSION_KEY, [])
    request.session[EXPORT_SESSION_KEY] = list(
        dict.fromkeys([*existing_ids, *reading_ids])
    )
    request.session.modified = True
    return created, updated, reading_ids


class IescoBillReadingListView(LoginRequiredMixin, ListView):
    template_name = "invoices/iesco_bill_dashboard.html"
    context_object_name = "meter_rows"
    paginate_by = 100

    def get_queryset(self):
        request = self.request
        reference_no = (request.GET.get("reference_no") or "").strip()
        bill_month = (request.GET.get("bill_month") or "").strip()
        payment_status = (request.GET.get("payment_status") or "").strip()
        reminder_review = request.GET.get("reminder_review") == "1"
        property_id = (request.GET.get("property") or "").strip()
        unit_id = (request.GET.get("unit") or "").strip()
        show_history = request.GET.get("show_history") == "1"
        period = (request.GET.get("period") or "this_month").strip()
        period_explicit = request.GET.get("period_explicit") == "1"

        sources = _visible_sources_for_filters(
            request.user,
            property_id=property_id,
            unit_id=unit_id,
            reference_no=reference_no,
        )
        references = {row["reference_no"] for row in sources}
        readings_qs = IescoBillReading.objects.filter(
            reference_no__in=references
        ).select_related("posted_invoice_item__invoice")
        readings, effective_bill_month = _select_dashboard_readings(
            readings_qs,
            references=references,
            bill_month=bill_month,
            period=period,
            period_explicit=period_explicit,
            history_mode=bool(show_history or unit_id or reference_no or period_explicit),
        )

        history_mode = bool(show_history or unit_id or reference_no or period_explicit)
        source_by_reference = {}
        for source in sources:
            source_by_reference.setdefault(source["reference_no"], source)

        rows = []
        if history_mode:
            for reading in readings:
                source = source_by_reference.get(reading.reference_no)
                if source is None:
                    continue
                if reminder_review and not source["is_active"]:
                    continue
                if not _payment_status_matches(reading, payment_status):
                    continue
                row = dict(source)
                row["reading"] = reading
                row["requires_payment"] = bool(
                    row["is_active"] and reading_requires_payment(reading)
                )
                rows.append(row)
        else:
            reading_by_reference = {reading.reference_no: reading for reading in readings}
            for source in sources:
                reading = reading_by_reference.get(source["reference_no"])
                if reminder_review and not source["is_active"]:
                    continue
                if payment_status and not _payment_status_matches(reading, payment_status):
                    continue
                row = dict(source)
                row["reading"] = reading
                row["requires_payment"] = bool(
                    row["is_active"] and reading and reading_requires_payment(reading)
                )
                rows.append(row)

        _attach_lease_context(rows)
        self.iesco_history_mode = history_mode
        self.iesco_effective_bill_month = effective_bill_month
        self.iesco_selected_period = period
        self.iesco_period_explicit = period_explicit
        return rows

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        properties = Property.objects.order_by("property_name")
        context["properties"] = restrict_queryset_to_properties(
            properties, self.request.user, ""
        )
        all_units = list(_valid_unit_meters(self.request.user))
        selected_property = (self.request.GET.get("property") or "").strip()
        selected_unit = (self.request.GET.get("unit") or "").strip()
        unit_filter_map = {}
        reference_options = []
        for unit in all_units:
            property_key = str(unit.property_id)
            unit_filter_map.setdefault(property_key, []).append(
                {
                    "id": str(unit.pk),
                    "label": unit.unit_number,
                    "property_name": unit.property.property_name,
                    "reference_no": unit.electric_meter_num,
                }
            )
            reference_options.append(
                {
                    "value": unit.electric_meter_num,
                    "label": f"{unit.electric_meter_num} — {unit.property.property_name} / {unit.unit_number}",
                }
            )
        standalone_meters = list(_visible_standalone_meters(self.request.user))
        reference_options.extend(
            {
                "value": meter.reference_no,
                "label": f"{meter.reference_no} — {meter.description or 'Standalone meter'}",
            }
            for meter in standalone_meters
        )
        reference_options.sort(key=lambda item: (item["label"], item["value"]))
        if selected_property and selected_property in unit_filter_map:
            unit_filter_options = unit_filter_map[selected_property]
        else:
            unit_filter_options = [
                {
                    "id": str(unit.pk),
                    "label": f"{unit.property.property_name} / {unit.unit_number}",
                    "property_name": unit.property.property_name,
                    "reference_no": unit.electric_meter_num,
                }
                for unit in all_units
            ]
        if selected_unit and not any(option["id"] == selected_unit for option in unit_filter_options):
            unit_filter_options = [
                option for option in unit_filter_options if option["id"] == selected_unit
            ] or unit_filter_options

        context["selected_property"] = selected_property
        context["selected_unit"] = selected_unit
        context["unit_filter_map"] = unit_filter_map
        context["unit_filter_options"] = unit_filter_options
        context["reference_options"] = reference_options
        context["standalone_form"] = IescoStandaloneMeterForm()
        context["can_manage"] = self.request.user.is_superuser or self.request.user.has_perm(
            "invoices.change_iescobillreading"
        )
        context["can_post"] = self.request.user.is_superuser or self.request.user.has_perm(
            "invoices.add_invoiceitem"
        )
        context["can_remind"] = context["can_manage"]
        context["can_add_standalone"] = has_all_property_access(
            self.request.user
        ) and (
            self.request.user.is_superuser
            or self.request.user.has_perm("invoices.add_iescostandalonemeter")
        )
        visible_references = set(
            _valid_unit_meters(self.request.user).values_list("electric_meter_num", flat=True)
        )
        visible_references.update(
            _visible_standalone_meters(self.request.user).values_list("reference_no", flat=True)
        )
        context["bill_month_options"] = _month_options_for_references(
            visible_references
        )
        context["period_options"] = IESCO_PERIOD_CHOICES
        context["selected_period"] = getattr(self, "iesco_selected_period", "this_month")
        context["period_explicit"] = getattr(self, "iesco_period_explicit", False)
        context["history_mode"] = getattr(self, "iesco_history_mode", False)
        context["effective_bill_month"] = getattr(self, "iesco_effective_bill_month", "")
        context["dashboard_totals"] = _dashboard_totals(self.object_list)
        context["has_iesco_preview"] = bool(
            self.request.session.get(PREVIEW_SESSION_KEY)
        )
        return context


class IescoBillReadingDetailView(LoginRequiredMixin, TemplateView):
    template_name = "invoices/iesco_bill_dashboard_detail.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        try:
            reference_no = validate_reference_no(kwargs["reference_no"])
        except ValidationError as exc:
            raise Http404("IESCO meter not found") from exc
        _ensure_reference_access(self.request.user, reference_no)
        units = list(
            _valid_unit_meters(self.request.user).filter(
                electric_meter_num=reference_no
            )
        )
        standalone = _visible_standalone_meters(self.request.user).filter(
            reference_no=reference_no
        ).first()
        readings = list(
            IescoBillReading.objects.filter(reference_no=reference_no)
            .select_related("posted_invoice_item__invoice")
            .order_by("-fetched_at", "-updated_at", "-id")
        )
        demo_mode = self.request.GET.get("demo") == "1"
        if demo_mode:
            readings = _demo_history_readings(
                reference_no, readings[0] if readings else None
            )
        detail_unit = units[0] if len(units) == 1 else None
        reading_rows = [
            {"reading": reading, "unit": detail_unit}
            for reading in readings
        ]
        _attach_lease_context(reading_rows)
        for row in reading_rows:
            reading = row["reading"]
            reading.occupancy_tenant_name = row["tenant_name"]
            reading.occupancy_lease_balance = row["lease_balance"]
            reading.occupancy_lease = row["lease"]
            reading.invoice_generated = row["invoice_generated"]
            reading.invoice_number = row["invoice_number"]
            reading.month_invoice = row["invoice"]
            reading.whatsapp_phone = row["whatsapp_phone"]
        for reading in readings:
            try:
                reading.posting_month_value = bill_month_start(
                    reading.bill_month
                ).strftime("%Y-%m")
            except ValidationError:
                reading.posting_month_value = timezone.localdate().strftime("%Y-%m")
        amounts = [reading.current_bill_amount for reading in readings if reading.current_bill_amount is not None]
        imported = [reading.import_units for reading in readings if reading.import_units is not None]
        total_amount = sum(amounts, Decimal("0"))
        total_imported = sum(imported, Decimal("0"))
        summary = {
            "bill_count": len(readings),
            "paid_count": sum(reading.current_month_paid is True for reading in readings),
            "total_billed": total_amount,
            "average_unit_rate": total_amount / total_imported if total_imported > 0 else None,
            "average_units": sum(imported, Decimal("0")) / len(imported) if imported else None,
        }
        context.update(
            {
                "reference_no": reference_no,
                "units": units,
                "unit": detail_unit,
                "standalone": standalone,
                "description": standalone.description if standalone else "",
                "readings": readings,
                "summary": summary,
                "demo_mode": demo_mode,
                "can_manage": self.request.user.is_superuser
                or self.request.user.has_perm("invoices.change_iescobillreading"),
                "can_post": self.request.user.is_superuser
                or self.request.user.has_perm("invoices.add_invoiceitem"),
            }
        )
        return context


@login_required
@require_POST
def standalone_meter_add(request):
    if not has_all_property_access(request.user):
        raise PermissionDenied("Standalone meters require access to all properties.")
    if not (
        request.user.is_superuser
        or request.user.has_perm("invoices.add_iescostandalonemeter")
    ):
        raise PermissionDenied
    form = IescoStandaloneMeterForm(request.POST)
    if form.is_valid():
        meter = form.save()
        messages.success(request, f"Standalone IESCO meter {meter.reference_no} added.")
    else:
        messages.error(
            request,
            "; ".join(
                str(message)
                for field_errors in form.errors.values()
                for message in field_errors
            ),
        )
    return redirect("invoices:iesco_bill_reading_list")


@login_required
@require_http_methods(["GET", "POST"])
def meter_edit(request, meter_kind, pk):
    _require_change_permission(request.user)
    unit_queryset = restrict_queryset_to_properties(
        Unit.objects.select_related("property").order_by(
            "property__property_name", "unit_number"
        ),
        request.user,
        "property",
    )
    current_unit = None
    current_standalone = None
    if meter_kind == "unit":
        current_unit = get_object_or_404(unit_queryset, pk=pk)
        title = f"{current_unit.property.property_name} / {current_unit.unit_number}"
        old_reference = current_unit.electric_meter_num
        initial = {
            "reference_no": old_reference,
            "unit": current_unit,
            "description": "",
            "phone": "",
            "is_active": current_unit.iesco_bill_active,
        }
    elif meter_kind == "standalone" and has_all_property_access(request.user):
        current_standalone = get_object_or_404(IescoStandaloneMeter, pk=pk)
        title = current_standalone.description
        old_reference = current_standalone.reference_no
        initial = {
            "reference_no": old_reference,
            "unit": None,
            "description": current_standalone.description,
            "phone": current_standalone.phone,
            "is_active": current_standalone.is_active,
        }
    else:
        raise Http404("IESCO meter not found")

    form = IescoMeterAssignmentForm(
        request.POST or None,
        initial=initial,
        unit_queryset=unit_queryset,
        current_unit=current_unit,
        current_standalone=current_standalone,
    )
    if request.method == "POST" and form.is_valid():
        new_reference = form.cleaned_data["reference_no"]
        target_unit = form.cleaned_data["unit"]
        description = (form.cleaned_data.get("description") or "").strip()
        phone = (form.cleaned_data.get("phone") or "").strip()
        with transaction.atomic():
            if target_unit:
                if current_unit and current_unit.pk != target_unit.pk:
                    current_unit.electric_meter_num = ""
                    current_unit.save(update_fields=["electric_meter_num"])
                target_unit.electric_meter_num = new_reference
                target_unit.iesco_bill_active = form.cleaned_data["is_active"]
                target_unit.save(
                    update_fields=["electric_meter_num", "iesco_bill_active"]
                )
                if current_standalone:
                    current_standalone.delete()
            else:
                if not has_all_property_access(request.user):
                    raise PermissionDenied(
                        "Standalone meters require access to all properties."
                    )
                if current_unit:
                    current_unit.electric_meter_num = ""
                    current_unit.save(update_fields=["electric_meter_num"])
                    current_standalone = IescoStandaloneMeter()
                current_standalone.reference_no = new_reference
                current_standalone.description = description
                current_standalone.phone = phone or None
                current_standalone.is_active = form.cleaned_data["is_active"]
                current_standalone.save()
        messages.success(request, f"IESCO meter updated to {new_reference}.")
        if old_reference != new_reference and IescoBillReading.objects.filter(
            reference_no=old_reference
        ).exists():
            messages.info(
                request,
                f"Existing bill history remains under its original reference {old_reference}.",
            )
        return redirect("invoices:iesco_bill_reading_list")
    return render(
        request,
        "invoices/iesco_meter_form.html",
        {
            "form": form,
            "title": title,
            "meter_kind": meter_kind,
            "meter_pk": pk,
        },
    )


@login_required
@require_POST
def meter_delete(request, meter_kind, pk):
    _require_change_permission(request.user)
    if meter_kind == "unit":
        queryset = restrict_queryset_to_properties(
            Unit.objects.select_related("property"), request.user, "property"
        )
        meter = get_object_or_404(queryset, pk=pk)
        reference_no = meter.electric_meter_num
        meter.electric_meter_num = ""
        meter.save(update_fields=["electric_meter_num"])
    elif meter_kind == "standalone" and has_all_property_access(request.user):
        meter = get_object_or_404(IescoStandaloneMeter, pk=pk)
        reference_no = meter.reference_no
        meter.delete()
    else:
        raise Http404("IESCO meter not found")
    messages.success(
        request,
        f"IESCO meter {reference_no} removed from the meter list. Saved bill history was preserved.",
    )
    return redirect("invoices:iesco_bill_reading_list")


@login_required
@require_GET
def pitc_bill_handoff(request, reference_no):
    reference_no = validate_reference_no(reference_no)
    _ensure_reference_access(request.user, reference_no)
    return render(
        request,
        "invoices/iesco_bill_external_redirect.html",
        {"reference_no": reference_no},
    )


@login_required
@require_POST
def fetch_one(request, reference_no):
    _require_change_permission(request.user)
    reference_no = validate_reference_no(reference_no)
    _ensure_reference_access(request.user, reference_no)
    description = _source_description(reference_no, request.user)
    try:
        payload = fetch_bill_payload(reference_no, description=description)
        created, updated, unused_ids = _save_payloads(request, [payload])
        action = "saved" if created else "updated"
        messages.success(
            request,
            f"IESCO bill for {reference_no} was {action} and the list is now current.",
        )
    except (ValidationError, requests.RequestException) as exc:
        messages.error(request, f"{reference_no}: {_fetch_failure_message(exc)}")
    return redirect("invoices:iesco_bill_reading_list")


@login_required
@require_POST
def fetch_all_active(request):
    _require_change_permission(request.user)
    property_id = (request.POST.get("property") or "").strip()
    sources = _active_fetch_sources(request.user, property_id)

    ajax_action = (request.POST.get("ajax_action") or "").strip()
    if ajax_action == "start":
        request.session[EXPORT_SESSION_KEY] = []
        request.session.modified = True
        return JsonResponse(
            {
                "sources": [
                    {
                        "reference_no": reference_no,
                        "label": f"{reference_no} — {description or 'Standalone meter'}",
                    }
                    for reference_no, description in sources.items()
                ],
                "total": len(sources),
                "list_url": reverse("invoices:iesco_bill_reading_list"),
            }
        )

    if ajax_action == "fetch":
        reference_no = (request.POST.get("reference_no") or "").strip()
        if reference_no not in sources:
            return JsonResponse(
                {"ok": False, "error": "This meter is not in the active fetch list."},
                status=400,
            )
        wait_until = _bulk_fetch_wait_until(reference_no)
        if wait_until:
            return JsonResponse(
                {
                    "ok": True,
                    "skipped": True,
                    "reference_no": reference_no,
                    "next_fetch": wait_until.strftime("%d %b %Y"),
                }
            )
        try:
            payload = fetch_bill_payload(
                reference_no, description=sources[reference_no]
            )
            created, updated, reading_ids = _save_payloads(request, [payload])
            existing_ids = request.session.get(EXPORT_SESSION_KEY, [])
            request.session[EXPORT_SESSION_KEY] = list(
                dict.fromkeys([*existing_ids, *reading_ids])
            )
            request.session.modified = True
            return JsonResponse(
                {
                    "ok": True,
                    "reference_no": reference_no,
                    "created": created,
                    "updated": updated,
                }
            )
        except (ValidationError, requests.RequestException) as exc:
            error = {"reference_no": reference_no, "error": _fetch_failure_message(exc)}
        except Exception as exc:
            logger.exception("IESCO AJAX bulk fetch failed for %s", reference_no)
            error = {"reference_no": reference_no, "error": _fetch_failure_message(exc)}
        return JsonResponse(
            {"ok": False, "reference_no": reference_no, "error": error["error"]}
        )

    errors = []
    skipped = 0
    created = 0
    updated = 0
    reading_ids = []
    for reference_no, description in sources.items():
        if _bulk_fetch_wait_until(reference_no):
            skipped += 1
            continue
        try:
            payload = fetch_bill_payload(reference_no, description=description)
            was_created, was_updated, saved_ids = _save_payloads(request, [payload])
            created += was_created
            updated += was_updated
            reading_ids.extend(saved_ids)
        except (ValidationError, requests.RequestException) as exc:
            errors.append(
                {"reference_no": reference_no, "error": _fetch_failure_message(exc)}
            )
        except Exception as exc:
            logger.exception("IESCO bulk fetch failed for %s", reference_no)
            errors.append(
                {"reference_no": reference_no, "error": _fetch_failure_message(exc)}
            )
    request.session[EXPORT_SESSION_KEY] = list(dict.fromkeys(reading_ids))
    request.session.modified = True
    if not created and not updated and not errors and skipped:
        messages.info(
            request,
            f"All {skipped} active IESCO bills are still within the fetch waiting period.",
        )
        return redirect("invoices:iesco_bill_reading_list")
    if created or updated:
        messages.success(
            request,
            f"IESCO bills saved: {created} new, {updated} updated, {skipped} current/cached.",
        )
    if errors:
        messages.warning(
            request,
            f"{len(errors)} IESCO bill(s) could not be fetched. {errors[0]['error']}",
        )
    return redirect("invoices:iesco_bill_reading_list")


@login_required
@require_POST
def import_csv(request):
    _require_change_permission(request.user)

    ajax_action = (request.POST.get("ajax_action") or "").strip()
    if ajax_action == "start":
        upload = request.FILES.get("file")
        if upload is None:
            return JsonResponse({"ok": False, "error": "Choose an IESCO CSV file to import."}, status=400)
        if upload.size > MAX_IMPORT_BYTES:
            return JsonResponse({"ok": False, "error": "IESCO import files cannot exceed 2 MB."}, status=400)
        try:
            text = upload.read().decode("utf-8-sig")
        except UnicodeDecodeError:
            return JsonResponse({"ok": False, "error": "The CSV file must use UTF-8 encoding."}, status=400)

        reader = csv.DictReader(io.StringIO(text))
        missing = {"reference_no", "bill_month"} - set(reader.fieldnames or ())
        if missing:
            return JsonResponse(
                {"ok": False, "error": "CSV is missing columns: " + ", ".join(sorted(missing))},
                status=400,
            )

        queued_rows = []
        display_rows = []
        row_errors = []
        for row_number, row in enumerate(reader, start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                row_errors.append(f"Import is limited to {MAX_IMPORT_ROWS} rows.")
                break
            safe_row = {str(key): (value or "") for key, value in row.items() if key is not None}
            try:
                payload = normalize_bill_payload(safe_row)
                try:
                    _ensure_reference_access(request.user, payload["reference_no"])
                except Http404:
                    if not (
                        has_all_property_access(request.user)
                        and payload.get("description")
                    ):
                        raise
            except (ValidationError, Http404) as exc:
                raw_reference = (safe_row.get("reference_no") or "").strip()
                row_errors.append(
                    f"Row {row_number}{f' ({raw_reference})' if raw_reference else ''}: {_validation_message(exc)}"
                )
                continue

            queued_rows.append({"row_number": row_number, "row": safe_row})
            display_rows.append(
                {
                    "row_number": row_number,
                    "reference_no": payload.get("reference_no") or "",
                    "bill_month": payload.get("bill_month") or "",
                    "consumer_name": payload.get("consumer_name") or "",
                    "description": payload.get("description") or "",
                }
            )

        if not queued_rows:
            request.session.pop(IMPORT_PROGRESS_SESSION_KEY, None)
            request.session.modified = True
            return JsonResponse(
                {
                    "ok": False,
                    "error": "The Bills CSV did not contain any valid bill rows.",
                    "row_errors": row_errors,
                },
                status=400,
            )

        request.session[IMPORT_PROGRESS_SESSION_KEY] = queued_rows
        request.session.modified = True
        return JsonResponse(
            {
                "ok": True,
                "total": len(display_rows),
                "rows": display_rows,
                "row_errors": row_errors,
            }
        )

    if ajax_action == "import":
        batch = request.session.get(IMPORT_PROGRESS_SESSION_KEY) or []
        try:
            row_index = int(request.POST.get("row_index", ""))
        except (TypeError, ValueError):
            return JsonResponse({"ok": False, "error": "Invalid import row."}, status=400)
        if row_index < 0 or row_index >= len(batch):
            return JsonResponse({"ok": False, "error": "Import row is no longer available. Restart the upload."}, status=400)

        stored = batch[row_index]
        try:
            payload = normalize_bill_payload(stored.get("row") or {})
            try:
                _ensure_reference_access(request.user, payload["reference_no"])
            except Http404:
                if not (
                    has_all_property_access(request.user)
                    and payload.get("description")
                ):
                    raise
            created, updated, _saved_ids = _save_payloads(request, [payload])
        except (ValidationError, Http404) as exc:
            return JsonResponse(
                {
                    "ok": False,
                    "reference_no": (stored.get("row") or {}).get("reference_no", ""),
                    "error": _validation_message(exc),
                }
            )
        except Exception as exc:
            logger.exception("IESCO Bills CSV AJAX import failed for row %s", stored.get("row_number"))
            return JsonResponse(
                {
                    "ok": False,
                    "reference_no": (stored.get("row") or {}).get("reference_no", ""),
                    "error": _validation_message(exc),
                }
            )

        return JsonResponse(
            {
                "ok": True,
                "reference_no": payload.get("reference_no") or "",
                "bill_month": payload.get("bill_month") or "—",
                "consumer_name": payload.get("consumer_name") or "—",
                "units": payload.get("units_display") or payload.get("units") or "—",
                "grand_total": payload.get("grand_total") or "—",
                "created": created,
                "updated": updated,
            }
        )

    if ajax_action == "finish":
        request.session.pop(IMPORT_PROGRESS_SESSION_KEY, None)
        request.session.modified = True
        return JsonResponse({"ok": True})

    upload = request.FILES.get("file")
    if upload is None:
        messages.error(request, "Choose an IESCO CSV file to import.")
        return redirect("invoices:iesco_bill_reading_list")
    if upload.size > MAX_IMPORT_BYTES:
        messages.error(request, "IESCO import files cannot exceed 2 MB.")
        return redirect("invoices:iesco_bill_reading_list")
    try:
        text = upload.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        messages.error(request, "The CSV file must use UTF-8 encoding.")
        return redirect("invoices:iesco_bill_reading_list")

    reader = csv.DictReader(io.StringIO(text))
    missing = {"reference_no", "bill_month"} - set(reader.fieldnames or ())
    if missing:
        messages.error(request, "CSV is missing columns: " + ", ".join(sorted(missing)))
        return redirect("invoices:iesco_bill_reading_list")

    payloads = []
    errors = []
    for index, row in enumerate(reader, start=2):
        if index > MAX_IMPORT_ROWS + 1:
            errors.append({"reference_no": "", "error": "Import is limited to 500 rows."})
            break
        try:
            payload = normalize_bill_payload(row)
            try:
                _ensure_reference_access(request.user, payload["reference_no"])
            except Http404:
                if not (
                    has_all_property_access(request.user)
                    and payload.get("description")
                ):
                    raise
            payloads.append(payload)
        except (ValidationError, Http404) as exc:
            errors.append(
                {
                    "reference_no": row.get("reference_no", ""),
                    "error": f"Row {index}: {_validation_message(exc)}",
                }
            )
    created, updated, unused_ids = _save_payloads(request, payloads) if payloads else (0, 0, [])
    if created or updated:
        messages.success(
            request,
            f"IESCO CSV imported: {created} new bill(s), {updated} updated.",
        )
    if errors:
        messages.warning(
            request,
            f"{len(errors)} CSV row(s) were skipped. " + errors[0]["error"],
        )
    if not payloads and not errors:
        messages.info(request, "The IESCO CSV did not contain any bill rows.")
    return redirect("invoices:iesco_bill_reading_list")


@login_required
@require_GET
def export_active_references(request):
    _require_change_permission(request.user)
    property_id = (request.GET.get("property") or "").strip()
    sources = _active_fetch_sources(request.user, property_id)
    if not sources:
        messages.error(request, "No active IESCO reference numbers are available for export.")
        return redirect("invoices:iesco_bill_reading_list")

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    filename = f"iesco-active-references-{timezone.localtime():%Y%m%d-%H%M%S}.csv"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write("\ufeff")
    writer = csv.DictWriter(response, fieldnames=["reference_no", "description"])
    writer.writeheader()
    for reference_no, description in sources.items():
        writer.writerow(
            {"reference_no": reference_no, "description": description or "IESCO meter"}
        )
    return response


@login_required
@require_POST
def fetch_reference_csv(request):
    _require_change_permission(request.user)
    if not has_all_property_access(request.user):
        raise PermissionDenied(
            "Importing the production reference list requires access to all properties."
        )

    ajax_action = (request.POST.get("ajax_action") or "").strip()
    if ajax_action == "start":
        upload = request.FILES.get("file")
        if upload is None:
            return JsonResponse(
                {"ok": False, "error": "Choose an active-reference CSV file."},
                status=400,
            )
        if upload.size > MAX_IMPORT_BYTES:
            return JsonResponse(
                {"ok": False, "error": "IESCO reference files cannot exceed 2 MB."},
                status=400,
            )
        try:
            text = upload.read().decode("utf-8-sig")
        except UnicodeDecodeError:
            return JsonResponse(
                {"ok": False, "error": "The reference CSV must use UTF-8 encoding."},
                status=400,
            )

        reader = csv.DictReader(io.StringIO(text))
        if "reference_no" not in set(reader.fieldnames or ()):
            return JsonResponse(
                {"ok": False, "error": "Reference CSV is missing the reference_no column."},
                status=400,
            )

        sources = {}
        row_errors = []
        for index, row in enumerate(reader, start=2):
            if index > MAX_IMPORT_ROWS + 1:
                row_errors.append(f"Reference import is limited to {MAX_IMPORT_ROWS} rows.")
                break
            raw_ref = (row.get("reference_no") or "").strip()
            if not raw_ref:
                continue
            try:
                reference_no = validate_reference_no(raw_ref)
                description = (row.get("description") or "").strip()
                sources[reference_no] = description or f"Imported IESCO meter {reference_no}"
            except ValidationError as exc:
                row_errors.append(f"Row {index} ({raw_ref}): {_clean_error_message(exc)}")

        if not sources:
            return JsonResponse(
                {"ok": False, "error": "The reference CSV did not contain any valid meter references."},
                status=400,
            )

        request.session[EXPORT_SESSION_KEY] = []
        request.session.modified = True

        meters = [
            {"reference_no": ref, "description": desc, "label": f"{ref} — {desc}"}
            for ref, desc in sources.items()
        ]
        return JsonResponse(
            {
                "ok": True,
                "total": len(meters),
                "meters": meters,
                "row_errors": row_errors,
            }
        )

    if ajax_action == "fetch":
        reference_no = (request.POST.get("reference_no") or "").strip()
        description = (request.POST.get("description") or "").strip() or f"Imported IESCO meter {reference_no}"
        if not reference_no:
            return JsonResponse(
                {"ok": False, "error": "Missing reference number."}, status=400
            )
        try:
            reference_no = validate_reference_no(reference_no)
        except ValidationError as exc:
            return JsonResponse(
                {"ok": False, "reference_no": reference_no, "error": _clean_error_message(exc)},
                status=400,
            )

        try:
            if not Unit.objects.filter(electric_meter_num=reference_no).exists():
                IescoStandaloneMeter.objects.update_or_create(
                    reference_no=reference_no,
                    defaults={"description": description, "is_active": True},
                )
            payload = fetch_bill_payload(reference_no, description=description)
            created, updated, saved_ids = _save_payloads(request, [payload])
            existing_ids = request.session.get(EXPORT_SESSION_KEY, [])
            request.session[EXPORT_SESSION_KEY] = list(
                dict.fromkeys([*existing_ids, *saved_ids])
            )
            request.session.modified = True
            return JsonResponse(
                {
                    "ok": True,
                    "reference_no": reference_no,
                    "created": created,
                    "updated": updated,
                    "bill_month": payload.get("bill_month") or "—",
                    "grand_total": payload.get("grand_total") or "—",
                    "units": payload.get("units_display") or payload.get("units") or "—",
                    "consumer_name": payload.get("consumer_name") or "—",
                }
            )
        except (ValidationError, requests.RequestException) as exc:
            return JsonResponse(
                {
                    "ok": False,
                    "reference_no": reference_no,
                    "error": _clean_error_message(exc),
                }
            )
        except Exception as exc:
            logger.exception("IESCO reference CSV AJAX fetch failed for %s", reference_no)
            return JsonResponse(
                {
                    "ok": False,
                    "reference_no": reference_no,
                    "error": _clean_error_message(exc),
                }
            )

    upload = request.FILES.get("file")
    if upload is None:
        messages.error(request, "Choose an active-reference CSV file.")
        return redirect("invoices:iesco_bill_reading_list")
    if upload.size > MAX_IMPORT_BYTES:
        messages.error(request, "IESCO reference files cannot exceed 2 MB.")
        return redirect("invoices:iesco_bill_reading_list")
    try:
        text = upload.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        messages.error(request, "The reference CSV must use UTF-8 encoding.")
        return redirect("invoices:iesco_bill_reading_list")

    reader = csv.DictReader(io.StringIO(text))
    if "reference_no" not in set(reader.fieldnames or ()):
        messages.error(request, "Reference CSV is missing the reference_no column.")
        return redirect("invoices:iesco_bill_reading_list")

    sources = {}
    errors = []
    for index, row in enumerate(reader, start=2):
        if index > MAX_IMPORT_ROWS + 1:
            errors.append("Reference import is limited to 500 rows.")
            break
        try:
            reference_no = validate_reference_no(row.get("reference_no"))
            description = (row.get("description") or "").strip()
            sources[reference_no] = description or f"Imported IESCO meter {reference_no}"
        except ValidationError as exc:
            errors.append(f"Row {index}: {_validation_message(exc)}")

    created = 0
    updated = 0
    reading_ids = []
    for reference_no, description in sources.items():
        try:
            if not Unit.objects.filter(electric_meter_num=reference_no).exists():
                IescoStandaloneMeter.objects.update_or_create(
                    reference_no=reference_no,
                    defaults={"description": description, "is_active": True},
                )
            payload = fetch_bill_payload(reference_no, description=description)
            was_created, was_updated, saved_ids = _save_payloads(request, [payload])
            created += was_created
            updated += was_updated
            reading_ids.extend(saved_ids)
        except (ValidationError, requests.RequestException) as exc:
            errors.append(f"{reference_no}: {_fetch_failure_message(exc)}")
        except Exception as exc:
            logger.exception("IESCO reference CSV fetch failed for %s", reference_no)
            errors.append(f"{reference_no}: {_fetch_failure_message(exc)}")

    request.session[EXPORT_SESSION_KEY] = list(dict.fromkeys(reading_ids))
    request.session.modified = True
    if created or updated:
        messages.success(
            request,
            f"Reference list processed and bills saved: {created} new, {updated} updated.",
        )
    if errors:
        messages.warning(
            request,
            f"{len(errors)} reference(s) could not be processed. {errors[0]}",
        )
    if not sources and not errors:
        messages.info(request, "The reference CSV did not contain any meters.")
    return redirect("invoices:iesco_bill_reading_list")


@login_required
@require_GET
def preview(request):
    _require_change_permission(request.user)
    batch = request.session.get(PREVIEW_SESSION_KEY)
    if not batch:
        messages.info(request, "There is no IESCO fetch or import waiting for review.")
        return redirect("invoices:iesco_bill_reading_list")
    return render(request, "invoices/iesco_bill_preview.html", {"batch": batch})


@login_required
@require_POST
def preview_save(request):
    _require_change_permission(request.user)
    batch = request.session.get(PREVIEW_SESSION_KEY)
    if not batch or not batch.get("payloads"):
        messages.error(request, "There are no valid IESCO bills to save.")
        return redirect("invoices:iesco_bill_preview")

    created = 0
    updated = 0
    reading_ids = []
    with transaction.atomic():
        for payload in batch["payloads"]:
            reference_no = payload["reference_no"]
            try:
                _ensure_reference_access(request.user, reference_no)
            except Http404:
                if not (
                    has_all_property_access(request.user)
                    and (payload.get("description") or "").strip()
                ):
                    raise
            reading, was_created = save_bill_payload(payload)
            reading_ids.append(reading.pk)
            created += int(was_created)
            updated += int(not was_created)
            description = (payload.get("description") or "").strip()
            if description and not Unit.objects.filter(
                electric_meter_num=reference_no
            ).exists():
                IescoStandaloneMeter.objects.update_or_create(
                    reference_no=reference_no,
                    defaults={"description": description, "is_active": True},
                )
    request.session.pop(PREVIEW_SESSION_KEY, None)
    request.session[EXPORT_SESSION_KEY] = reading_ids
    request.session.modified = True
    messages.success(
        request,
        f"IESCO bills saved: {created} new, {updated} updated. CSV export is ready.",
    )
    return redirect("invoices:iesco_bill_reading_list")


def _reading_export_payload(reading, user):
    return {
        "reference_no": reading.reference_no,
        "fetched_at": reading.fetched_at.isoformat() if reading.fetched_at else None,
        "consumer_id": reading.consumer_id,
        "consumer_name": reading.consumer_name,
        "address": reading.address,
        "tariff_category": reading.tariff_category,
        "meter_type": reading.meter_type,
        "meter_readings": reading.meter_readings,
        "units": reading.units,
        "bill_month": reading.bill_month,
        "reading_date": reading.reading_date,
        "issue_date": reading.issue_date,
        "due_date": reading.due_date,
        "current_bill": reading.current_bill,
        "arrears": reading.arrears,
        "grand_total": reading.grand_total,
        "amount_paid": reading.amount_paid,
        "payment_date": reading.payment_date,
        "bill_history": reading.bill_history,
        "current_month_paid": reading.current_month_paid,
        "description": _source_description(reading.reference_no, user),
    }


def _export_readings_for_request(request, *, reading_id="", batch_only=False):
    property_id = (request.GET.get("property") or "").strip()
    unit_id = (request.GET.get("unit") or "").strip()
    reference_no = (request.GET.get("reference_no") or "").strip()
    bill_month = (request.GET.get("bill_month") or "").strip()
    payment_status = (request.GET.get("payment_status") or "").strip()
    show_history = request.GET.get("show_history") == "1"
    period = (request.GET.get("period") or "this_month").strip()
    period_explicit = request.GET.get("period_explicit") == "1"

    sources = _visible_sources_for_filters(
        request.user,
        property_id=property_id,
        unit_id=unit_id,
        reference_no=reference_no,
    )
    visible_references = {row["reference_no"] for row in sources}
    readings = IescoBillReading.objects.filter(
        reference_no__in=visible_references
    ).select_related("posted_invoice_item__invoice")

    if reading_id and str(reading_id).isdigit():
        selected = list(readings.filter(pk=int(reading_id)))
    elif batch_only:
        batch_ids = [
            int(value)
            for value in request.session.get(EXPORT_SESSION_KEY, [])
            if str(value).isdigit()
        ]
        if not batch_ids:
            return [], sources
        selected = list(
            readings.filter(pk__in=batch_ids).order_by(
                "-fetched_at", "-updated_at", "-id"
            )
        )
    else:
        history_mode = bool(
            show_history or unit_id or reference_no or period_explicit
        )
        selected, _effective_bill_month = _select_dashboard_readings(
            readings,
            references=visible_references,
            bill_month=bill_month,
            period=period,
            period_explicit=period_explicit,
            history_mode=history_mode,
        )

    if payment_status:
        selected = [
            reading
            for reading in selected
            if _payment_status_matches(reading, payment_status)
        ]
    return selected, sources


@login_required
@require_GET
def export_last_csv(request):
    batch_only = (request.GET.get("batch") or "").strip() == "1"
    visible, _sources = _export_readings_for_request(
        request, batch_only=batch_only
    )
    if not visible:
        if batch_only:
            messages.error(request, "No newly fetched IESCO bills are ready for download.")
        else:
            messages.error(request, "No saved IESCO bills are available for export.")
        return redirect("invoices:iesco_bill_reading_list")

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    filename_prefix = "iesco-bills-fetched" if batch_only else "iesco-bills"
    filename = f"{filename_prefix}-{timezone.localtime():%Y%m%d-%H%M%S}.csv"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write("\ufeff")
    writer = csv.DictWriter(response, fieldnames=IESCO_EXPORT_FIELDS)
    writer.writeheader()
    for reading in visible:
        writer.writerow(payload_to_csv_row(_reading_export_payload(reading, request.user)))
    return response


def _formatted_register_value(reading, field):
    registers = reading.register_display_rows
    if not registers:
        return "—"
    if len(registers) == 1:
        return str(registers[0].get(field) or "—")
    return " / ".join(
        f"{register['short_label']}: {register.get(field) or '—'}"
        for register in registers
    )


def _formatted_units(reading):
    if reading.has_export_registers:
        return (
            f"E:{reading.export_units_display} / I:{reading.import_units_display} / "
            f"N:{reading.net_units_display}"
        )
    return reading.units_display


def _formatted_export_rows(request):
    reading_id = (request.GET.get("reading_id") or "").strip()
    ordered, sources = _export_readings_for_request(
        request,
        reading_id=reading_id,
    )
    source_by_reference = {}
    for source in sources:
        source_by_reference.setdefault(source["reference_no"], source)

    context_rows = []
    for reading in ordered:
        source = source_by_reference.get(reading.reference_no, {})
        context_rows.append(
            {
                "reading": reading,
                "unit": source.get("unit"),
                "standalone": source.get("standalone"),
                "is_active": bool(source.get("is_active")),
            }
        )
    _attach_lease_context(context_rows)

    rows = []
    for serial, context in enumerate(context_rows, 1):
        reading = context["reading"]
        unit = context["unit"]
        standalone = context["standalone"]
        property_or_description = (
            unit.property.property_name if unit else (standalone.description if standalone else "")
        )
        unit_number = unit.unit_number if unit else "Standalone meter"
        previous = _formatted_register_value(reading, "previous")
        current = _formatted_register_value(reading, "present")
        active = "Active" if context["is_active"] else "Inactive"
        tenant = context["tenant_name"] or "—"
        balance = (
            f"Rs. {context['lease_balance']:,.0f}"
            if context["lease_balance"] is not None
            else "—"
        )
        invoice_number = context["invoice_number"] or "—"
        rows.append(
            {
                "serial": serial,
                "property_or_description": property_or_description or "—",
                "unit": unit_number,
                "property_unit": f"{property_or_description or '—'}\n{unit_number}",
                "active": active,
                "reference_no": reading.reference_no,
                "meter_number": reading.meter_number_display,
                "reference_meter": f"{reading.reference_no}\n{reading.meter_number_display}",
                "bill_month": reading.bill_month,
                "bill_active": f"{reading.bill_month}\n{active}",
                "consumer": reading.consumer_name or reading.consumer_id or "—",
                "tenant": tenant,
                "consumer_tenant": f"{reading.consumer_name or reading.consumer_id or '—'}\n{tenant}",
                "previous_reading": previous,
                "current_reading": current,
                "readings": f"Previous: {previous}\nCurrent: {current}",
                "units": _formatted_units(reading),
                "per_unit": reading.per_unit_rate_display,
                "arrears": reading.arrears_display,
                "current_bill": reading.current_bill_display,
                "charges": f"Arrears: {reading.arrears_display}\nCurrent: {reading.current_bill_display}",
                "grand_total": reading.grand_total_display,
                "due_date": reading.due_date or "—",
                "payment_status": reading.payment_status_display,
                "due_payment": f"{reading.due_date or '—'}\n{reading.payment_status_display}",
                "reading_date": reading.reading_date or "—",
                "last_update": timezone.localtime(reading.updated_at).strftime("%Y-%m-%d %H:%M"),
                "lease_balance": balance,
                "update_balance": f"{timezone.localtime(reading.updated_at).strftime('%Y-%m-%d %H:%M')}\nBalance: {balance}",
                "invoice_number": invoice_number,
                "is_summary": False,
            }
        )

    if rows and not reading_id:
        totals = _dashboard_totals(context_rows)
        summary = {key: "" for key in rows[0].keys()}
        summary.update(
            {
                "serial": "",
                "property_or_description": "TOTAL",
                "unit": "",
                "property_unit": "TOTAL",
                "bill_month": "",
                "bill_active": "",
                "reference_meter": "",
                "consumer_tenant": "",
                "readings": "",
                "units": totals["units_display"],
                "grand_total": totals["grand_total_display"],
                "is_summary": True,
            }
        )
        rows.append(summary)
    return rows


def _iesco_export_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = [
        "S.N", "Property / Description", "Unit", "Active", "Reference #", "Meter #",
        "Bill Month", "Consumer", "Occupant", "Previous Reading", "Current Reading",
        "Units", "Per Unit", "Arrears", "Current Bill", "Grand Total",
        "Payment Status", "Due Date", "Reading Date", "Last Update", "Lease Balance",
        "Invoice Number",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "IESCO Bills"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="198754")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    keys = [
        "serial", "property_or_description", "unit", "active", "reference_no", "meter_number",
        "bill_month", "consumer", "tenant", "previous_reading", "current_reading", "units",
        "per_unit", "arrears", "current_bill", "grand_total", "payment_status", "due_date",
        "reading_date", "last_update", "lease_balance", "invoice_number",
    ]
    for row in rows:
        sheet.append([row[key] for key in keys])
        if row.get("is_summary"):
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E9ECEF")
    widths = [7, 24, 18, 10, 18, 18, 13, 24, 24, 24, 24, 22, 13, 14, 15, 15, 16, 15, 15, 18, 16, 20]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _image_font(size, *, bold=False):
    from PIL import ImageFont

    names = ["arialbd.ttf" if bold else "arial.ttf", "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _iesco_export_jpg(rows, *, single=False):
    from PIL import Image, ImageDraw

    if single:
        row = rows[0]
        image = Image.new("RGB", (1080, 1080), "white")
        draw = ImageDraw.Draw(image)
        draw.rounded_rectangle((45, 45, 1035, 1035), radius=30, fill="#f8fafc", outline="#198754", width=5)
        draw.rectangle((45, 45, 1035, 205), fill="#198754")
        draw.text((90, 85), "IESCO BILL SUMMARY", fill="white", font=_image_font(48, bold=True))
        labels = [
            ("Property / Unit", row["property_unit"]), ("Month / Active", row["bill_active"]),
            ("Reference / Meter", row["reference_meter"]), ("Consumer / Occupant", row["consumer_tenant"]),
            ("Previous / Current", row["readings"]), ("Units", row["units"]),
            ("Per Unit", row["per_unit"]), ("Arrears / Current", row["charges"]),
            ("Grand Total", row["grand_total"]), ("Due / Payment", row["due_payment"]),
            ("Updated / Balance", row["update_balance"]),
        ]
        y = 225
        for label, value in labels:
            draw.text((90, y), label, fill="#64748b", font=_image_font(20, bold=True))
            draw.multiline_text((365, y), str(value), fill="#0f172a", font=_image_font(21, bold=True), spacing=3)
            draw.line((90, y + 62, 990, y + 62), fill="#dbe3ea", width=2)
            y += 68
    else:
        widths = [210, 115, 185, 205, 210, 145, 115, 165, 140, 150, 185]
        headers = ["Property / Unit", "Month / Active", "Reference / Meter", "Consumer / Occupant", "Previous / Current", "Units", "Per Unit", "Arrears / Current", "Grand Total", "Due / Payment", "Updated / Balance"]
        row_height = 76
        image = Image.new("RGB", (sum(widths) + 40, 100 + row_height * (len(rows) + 1)), "white")
        draw = ImageDraw.Draw(image)
        draw.text((20, 18), "IESCO BILL READINGS", fill="#0f172a", font=_image_font(32, bold=True))
        y = 78
        x = 20
        for header, width in zip(headers, widths):
            draw.rectangle((x, y, x + width, y + row_height), fill="#198754", outline="#d1d5db")
            draw.multiline_text((x + 7, y + 16), header, fill="white", font=_image_font(16, bold=True), spacing=2)
            x += width
        for row in rows:
            y += row_height
            x = 20
            values = [row["property_unit"], row["bill_active"], row["reference_meter"], row["consumer_tenant"], row["readings"], row["units"], row["per_unit"], row["charges"], row["grand_total"], row["due_payment"], row["update_balance"]]
            for value, width in zip(values, widths):
                draw.rectangle((x, y, x + width, y + row_height), fill="#ffffff", outline="#d1d5db")
                draw.multiline_text((x + 7, y + 12), str(value)[:70], fill="#111827", font=_image_font(14), spacing=2)
                x += width
    stream = io.BytesIO()
    image.save(stream, format="JPEG", quality=92, optimize=True)
    return stream.getvalue()


@login_required
@require_GET
def export_formatted(request, export_format):
    if export_format not in {"xlsx", "pdf", "jpg"}:
        raise Http404
    rows = _formatted_export_rows(request)
    if not rows:
        messages.error(request, "No saved IESCO bills are available for export.")
        return redirect("invoices:iesco_bill_reading_list")
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M%S")
    single = len(rows) == 1
    if export_format == "xlsx":
        content = _iesco_export_xlsx(rows)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif export_format == "jpg":
        content = _iesco_export_jpg(rows, single=single)
        content_type = "image/jpeg"
    else:
        from weasyprint import HTML

        html = render_to_string("invoices/iesco_bill_export_pdf.html", {"rows": rows, "single": single})
        content = HTML(string=html, base_url=request.build_absolute_uri("/")).write_pdf()
        content_type = "application/pdf"
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="iesco-bills-{stamp}.{export_format}"'
    return response


def _reminder_preview_row(reading):
    recipient = reminder_recipient(reading)
    return {
        "reading": reading,
        "recipient": recipient,
        "message": build_iesco_reminder_message(reading, recipient),
    }


@login_required
@require_POST
def reminders_send(request):
    _require_change_permission(request.user)
    raw_ids = list(dict.fromkeys(request.POST.getlist("reading_id")))[:100]
    readings = list(IescoBillReading.objects.filter(pk__in=raw_ids))
    for reading in readings:
        _ensure_reference_access(request.user, reading.reference_no)
    is_single_status_message = len(raw_ids) == 1 and len(readings) == 1
    eligible = (
        readings
        if is_single_status_message
        else [reading for reading in readings if reading_requires_payment(reading)]
    )
    if request.POST.get("confirm") != "1":
        preview_rows = [_reminder_preview_row(reading) for reading in eligible]
        return render(
            request,
            "invoices/iesco_reminder_preview.html",
            {
                "rows": preview_rows,
                "has_sendable": any(
                    row["recipient"]["phone"] for row in preview_rows
                ),
            },
        )
    sent = 0
    skipped = 0
    first_error = ""
    for reading in eligible:
        result = send_iesco_reminder(
            reading.pk,
            user=request.user,
            allow_any_status=is_single_status_message,
        )
        if result.get("ok"):
            sent += 1
        else:
            skipped += 1
            first_error = first_error or result.get("reason", "")
    if sent:
        messages.success(request, f"{sent} IESCO WhatsApp message(s) sent through Meta.")
    if skipped:
        messages.warning(request, f"{skipped} reminder(s) skipped. {first_error}")
    return redirect("invoices:iesco_bill_reading_list")


@login_required
@require_POST
def upload_bill_pdf(request, pk):
    _require_change_permission(request.user)
    reading = get_object_or_404(IescoBillReading, pk=pk)
    _ensure_reference_access(request.user, reading.reference_no)
    upload = request.FILES.get("bill_pdf")
    if upload is None:
        messages.error(request, "Choose a PDF bill to upload.")
    elif reading.bill_pdf:
        messages.error(request, "A PDF is already stored for this monthly bill.")
    elif upload.size > 10 * 1024 * 1024:
        messages.error(request, "IESCO PDF files cannot exceed 10 MB.")
    else:
        header = upload.read(5)
        upload.seek(0)
        if header != b"%PDF-":
            messages.error(request, "The selected file is not a valid PDF.")
        else:
            reading.bill_pdf.save(upload.name, upload, save=False)
            reading.pdf_uploaded_at = timezone.now()
            reading.save(update_fields=["bill_pdf", "pdf_uploaded_at", "updated_at"])
            messages.success(request, "IESCO bill PDF uploaded.")
    return redirect(
        reverse("invoices:iesco_bill_reading_detail", args=[reading.reference_no])
    )


@login_required
@require_POST
def post_to_invoice(request, pk):
    if not (
        request.user.is_superuser or request.user.has_perm("invoices.add_invoiceitem")
    ):
        raise PermissionDenied
    reading = IescoBillReading.objects.filter(pk=pk).first()
    if reading is None:
        raise Http404
    _ensure_reference_access(request.user, reading.reference_no)
    try:
        posting_month = bill_month_start(reading.bill_month)
        invoice, _item, action = post_reading_to_invoice(
            reading.pk,
            posting_month,
            user=request.user,
            confirm_amount_change=request.POST.get("confirm_amount_change") == "1",
        )
    except IescoInvoiceAmountChangeRequired as exc:
        return render(
            request,
            "invoices/iesco_invoice_amount_changes.html",
            {
                "changes": [_invoice_change_row(exc)],
                "confirm_url": reverse(
                    "invoices:iesco_bill_post_to_invoice", args=[reading.pk]
                ),
                "return_url": reverse(
                    "invoices:iesco_bill_reading_detail",
                    args=[reading.reference_no],
                ),
                "title": "Confirm IESCO invoice update",
            },
        )
    except ValidationError as exc:
        messages.error(request, _validation_message(exc))
    else:
        success_text = {
            "created": f"IESCO bill posted to invoice {invoice.invoice_number}.",
            "updated": f"IESCO bill updated on invoice {invoice.invoice_number}.",
            "unchanged": f"IESCO bill already matches invoice {invoice.invoice_number}.",
        }[action]
        messages.success(request, success_text)
    if request.POST.get("return_to") == "list":
        return redirect("invoices:iesco_bill_reading_list")
    return redirect("invoices:iesco_bill_reading_detail", reference_no=reading.reference_no)


def _invoice_change_row(change):
    unit = (
        Unit.objects.select_related("property")
        .filter(electric_meter_num=change.reading.reference_no)
        .first()
    )
    return {
        "reading_id": change.reading.pk,
        "reference_no": change.reading.reference_no,
        "bill_month": change.reading.bill_month,
        "unit_label": (
            f"{unit.property.property_name} / {unit.unit_number}"
            if unit
            else "Unassigned meter"
        ),
        "invoice_number": change.invoice.invoice_number,
        "old_amount": change.old_amount,
        "new_amount": change.new_amount,
        "difference": change.new_amount - change.old_amount,
    }


@login_required
@require_POST
def make_invoices_bulk(request):
    if not (
        request.user.is_superuser or request.user.has_perm("invoices.add_invoiceitem")
    ):
        raise PermissionDenied

    raw_ids = list(dict.fromkeys(request.POST.getlist("reading_id")))
    if len(raw_ids) > 100:
        messages.error(request, "Bulk invoice creation is limited to 100 readings.")
        return redirect("invoices:iesco_bill_reading_list")
    try:
        reading_ids = [int(value) for value in raw_ids]
    except (TypeError, ValueError):
        messages.error(request, "The selected IESCO readings are invalid.")
        return redirect("invoices:iesco_bill_reading_list")
    if not reading_ids:
        messages.info(request, "No invoice-ready IESCO readings are visible.")
        return redirect("invoices:iesco_bill_reading_list")

    readings = {
        reading.pk: reading
        for reading in IescoBillReading.objects.filter(pk__in=reading_ids)
    }
    confirm_changes = request.POST.get("confirm_amount_change") == "1"
    counts = {"created": 0, "updated": 0, "unchanged": 0}
    changes = []
    errors = []
    for reading_id in reading_ids:
        reading = readings.get(reading_id)
        if reading is None:
            errors.append(f"Reading {reading_id} was not found.")
            continue
        try:
            _ensure_reference_access(request.user, reading.reference_no)
            posting_month = bill_month_start(reading.bill_month)
            _invoice, _item, action = post_reading_to_invoice(
                reading.pk,
                posting_month,
                user=request.user,
                confirm_amount_change=confirm_changes,
            )
            counts[action] += 1
        except IescoInvoiceAmountChangeRequired as exc:
            changes.append(_invoice_change_row(exc))
        except (Http404, ValidationError) as exc:
            errors.append(f"{reading.reference_no}: {_validation_message(exc)}")

    if changes:
        return render(
            request,
            "invoices/iesco_invoice_amount_changes.html",
            {
                "changes": changes,
                "errors": errors,
                "counts": counts,
                "confirm_url": reverse("invoices:iesco_bill_make_invoices_bulk"),
                "return_url": reverse("invoices:iesco_bill_reading_list"),
                "title": "Confirm changed IESCO invoice amounts",
            },
        )

    if any(counts.values()):
        messages.success(
            request,
            "IESCO invoices processed: "
            f"{counts['created']} created, {counts['updated']} updated, "
            f"{counts['unchanged']} unchanged.",
        )
    if errors:
        messages.warning(request, f"{len(errors)} reading(s) were skipped. {errors[0]}")
    return redirect("invoices:iesco_bill_reading_list")
