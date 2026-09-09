from datetime import date, timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image

from invoices.models import IescoBillReading, IescoStandaloneMeter
from invoices.services_iesco_reminders import reminder_recipient
from invoices.views_iesco import _bulk_fetch_wait_until
from leases.models import Lease
from properties.models import Property, Unit
from tenants.models import Tenant


class IescoReminderAndExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            username="iesco-reminder-admin",
            email="iesco-reminder@example.com",
            password="test-pass",
        )
        self.client.force_login(self.user)
        self.property = Property.objects.create(
            property_name="Reminder Property",
            owner_name="Property Owner",
            owner_phone="03001112222",
            owner_cnic="12345-1234567-1",
            type="residential",
            property_type="apartment",
            total_units=1,
        )
        self.unit = Unit.objects.create(
            property=self.property,
            unit_number="A-1",
            electric_meter_num="17146151548001",
            iesco_bill_active=True,
        )
        self.tenant = Tenant.objects.create(
            first_name="Test",
            last_name="Tenant",
            phone="03002223333",
            cnic="12345-1234567-2",
        )
        self.lease = Lease.objects.create(
            tenant=self.tenant,
            unit=self.unit,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            monthly_rent="25000",
            status="active",
        )

    def reading(self, **overrides):
        values = {
            "reference_no": self.unit.electric_meter_num,
            "bill_month": "SEP 26",
            "units": "100",
            "current_bill": "2,500",
            "arrears": "0",
            "grand_total": "2,500",
            "due_date": "09 SEP 26",
            "current_month_paid": False,
            "fetched_at": timezone.now(),
        }
        values.update(overrides)
        return IescoBillReading.objects.create(**values)

    def test_negative_grand_total_is_not_payment_due_or_unpaid_filtered(self):
        reading = self.reading(grand_total="-1,422", current_bill="684", arrears="-2,106")

        self.assertEqual(reading.payment_status_display, "No payment due")
        response = self.client.get(
            reverse("invoices:iesco_bill_reading_list"),
            {"payment_status": "unpaid"},
        )
        self.assertNotContains(response, reading.reference_no)

    def test_recipient_prefers_tenant_then_owner_and_standalone_phone(self):
        reading = self.reading()
        self.assertEqual(reminder_recipient(reading)["phone"], self.tenant.phone)

        self.tenant.phone = ""
        self.tenant.save(update_fields=["phone"])
        self.assertEqual(reminder_recipient(reading)["phone"], self.property.owner_phone)

        standalone = IescoStandaloneMeter.objects.create(
            reference_no="17146151548002",
            description="Standalone Office",
            phone="03003334444",
        )
        standalone_reading = self.reading(
            reference_no=standalone.reference_no,
            bill_month="AUG 26",
        )
        self.assertEqual(reminder_recipient(standalone_reading)["phone"], standalone.phone)

    @patch("whatsapp.services.whatsapp.WhatsAppService.send_image_bytes", return_value={"ok": True})
    def test_single_reminder_requires_review_then_sends_once(self, send_image):
        reading = self.reading()
        url = reverse("invoices:iesco_bill_reminders_send")

        preview = self.client.post(url, {"reading_id": reading.pk})
        self.assertContains(preview, "Review IESCO WhatsApp Messages")
        self.assertContains(preview, self.tenant.phone)
        self.assertContains(preview, "Unpaid")
        send_image.assert_not_called()

        sent = self.client.post(
            url,
            {"reading_id": reading.pk, "confirm": "1"},
            follow=True,
        )
        self.assertContains(sent, "1 IESCO WhatsApp message")
        body = send_image.call_args.kwargs["caption"]
        self.assertIn(reading.reference_no, body)
        self.assertIn("Payment status: Unpaid", body)
        self.assertIn("bill.pitc.com.pk", body)
        self.assertTrue(send_image.call_args.args[1].startswith(b"\xff\xd8"))
        reading.refresh_from_db()
        self.assertIsNotNone(reading.reminder_sent_at)

        self.client.post(url, {"reading_id": reading.pk, "confirm": "1"})
        self.assertEqual(send_image.call_count, 1)

    @patch("whatsapp.services.whatsapp.WhatsAppService.send_image_bytes", return_value={"ok": True})
    def test_paid_bill_status_can_be_reviewed_and_sent(self, send_image):
        reading = self.reading(current_month_paid=True)
        url = reverse("invoices:iesco_bill_reminders_send")

        listing = self.client.get(reverse("invoices:iesco_bill_reading_list"))
        self.assertContains(listing, "Send bill status by WhatsApp")

        preview = self.client.post(url, {"reading_id": reading.pk})
        self.assertContains(preview, "Paid")

        sent = self.client.post(
            url,
            {"reading_id": reading.pk, "confirm": "1"},
            follow=True,
        )

        self.assertContains(sent, "1 IESCO WhatsApp message")
        self.assertIn(
            "Payment status: Paid", send_image.call_args.kwargs["caption"]
        )

    def test_excel_keeps_separate_fields_and_jpg_pdf_use_compact_layout(self):
        reading = self.reading()
        query = f"?reading_id={reading.pk}"

        excel = self.client.get(
            reverse("invoices:iesco_bill_export_formatted", args=["xlsx"]) + query
        )
        self.assertEqual(excel.status_code, 200)
        workbook = load_workbook(BytesIO(excel.content))
        headers = [cell.value for cell in workbook.active[1]]
        self.assertEqual(headers[0], "S.N")
        self.assertIn("Reference #", headers)
        self.assertIn("Meter #", headers)
        self.assertIn("Previous Reading", headers)
        self.assertIn("Current Reading", headers)
        self.assertIn("Arrears", headers)
        self.assertIn("Current Bill", headers)
        self.assertEqual(
            workbook.active.cell(row=2, column=headers.index("Reference #") + 1).value,
            reading.reference_no,
        )

        jpg = self.client.get(
            reverse("invoices:iesco_bill_export_formatted", args=["jpg"]) + query
        )
        self.assertEqual(jpg.status_code, 200)
        image = Image.open(BytesIO(jpg.content))
        self.assertEqual(image.size, (1080, 1080))

        pdf = self.client.get(
            reverse("invoices:iesco_bill_export_formatted", args=["pdf"]) + query
        )
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_unpaid_status_refreshes_after_three_days_but_paid_uses_long_cache(self):
        reading = self.reading()
        reading.fetched_at = timezone.now() - timedelta(days=4)
        reading.save(update_fields=["fetched_at"])
        self.assertIsNone(_bulk_fetch_wait_until(reading.reference_no))

        reading.current_month_paid = True
        reading.fetched_at = timezone.now()
        reading.reading_date = ""
        reading.save(update_fields=["current_month_paid", "fetched_at", "reading_date"])
        wait_until = _bulk_fetch_wait_until(reading.reference_no)
        self.assertGreaterEqual(wait_until, timezone.localdate() + timedelta(days=19))
