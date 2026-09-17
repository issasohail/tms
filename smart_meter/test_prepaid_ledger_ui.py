from decimal import Decimal
from datetime import datetime, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from properties.models import Property, Unit
from smart_meter.models import (
    LiveReading,
    Meter,
    MeterCommand,
    MeterPrepaidPilot,
    MeterPrepaidRecharge,
    MeterReading,
)


class PrepaidLedgerUITests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            username="prepaid-ledger-admin",
            password="pass",
            email="ledger@example.com",
        )
        self.client.force_login(self.user)
        property_obj = Property.objects.create(
            property_name="Ledger Property",
            owner_name="Owner",
            owner_cnic="1234512345673",
            type="apartment",
            property_type="apartment",
            total_units=1,
        )
        unit = Unit.objects.create(property=property_obj, unit_number="L-1")
        self.meter = Meter.objects.create(
            meter_number="260305519991",
            unit=unit,
            billing_mode="prepaid_pilot",
            tariff_capability="single_rate",
            unit_rate=Decimal("50.0000"),
        )
        MeterPrepaidPilot.objects.create(meter=self.meter, status="active_test")
        LiveReading.objects.create(
            meter=self.meter,
            balance=Decimal("95.00"),
            total_energy=Decimal("10.100"),
            voltage_a=Decimal("230.0"),
            current_a=Decimal("1.250"),
        )
        MeterReading.objects.create(
            meter=self.meter,
            balance=Decimal("100.00"),
            total_energy=Decimal("10.000"),
            unit_rate=Decimal("50.0000"),
        )
        MeterReading.objects.create(
            meter=self.meter,
            balance=Decimal("95.00"),
            total_energy=Decimal("10.100"),
            unit_rate=Decimal("50.0000"),
        )

    def test_ledger_shows_balance_usage_rate_and_money_form(self):
        response = self.client.get(
            reverse("smart_meter:prepaid_meter_ledger", args=[self.meter.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Meter Balance Ledger")
        self.assertContains(response, "Estimated usage charge")
        self.assertContains(response, "Queue once")
        self.assertContains(response, "95.00")
        self.assertTrue(response.context["has_collapsible_readings"])
        self.assertTrue(response.context["readings"][0].is_daily_latest)
        self.assertFalse(response.context["readings"][1].is_daily_latest)
        self.assertContains(response, 'id="toggleReadingDays"')
        self.assertContains(response, "ledger-reading-extra d-none")
        self.assertContains(response, "ledger-mobile-detail d-none")

    def test_prepaid_controls_show_operational_columns_and_ledger_link(self):
        response = self.client.get(reverse("smart_meter:prepaid_controls"))

        self.assertEqual(response.status_code, 200)
        for label in ("S/N", "Online", "Balance", "Unit rate", "Voltage", "Current"):
            self.assertContains(response, label)
        self.assertContains(
            response,
            reverse("smart_meter:prepaid_meter_ledger", args=[self.meter.pk]),
        )

    def test_ledger_topup_queues_one_guarded_money_command(self):
        response = self.client.post(
            reverse("smart_meter:prepaid_meter_ledger", args=[self.meter.pk]),
            {
                "operation": "recharge",
                "amount": "5.00",
                "reason": "UI ledger test",
                "confirm_meter_number": self.meter.meter_number,
            },
        )

        self.assertRedirects(
            response,
            reverse("smart_meter:prepaid_meter_ledger", args=[self.meter.pk]),
        )
        self.assertEqual(MeterPrepaidRecharge.objects.count(), 1)
        command = MeterCommand.objects.get(command_type="prepaid_recharge")
        self.assertEqual(command.max_attempts, 1)
        self.assertEqual(command.expect_di, "070102FF")

    def test_meter_and_live_lists_link_balance_to_ledger(self):
        ledger_url = f'{reverse("smart_meter:meter_detail", args=[self.meter.pk])}?tab=ledger'
        meter_response = self.client.get(reverse("smart_meter:meter_list"))
        live_response = self.client.get(reverse("smart_meter:smart_meter_live_custom"))
        reading_response = self.client.get(reverse("smart_meter:reading_list"))
        controls_response = self.client.get(reverse("smart_meter:prepaid_controls"))
        detail_response = self.client.get(ledger_url)

        for response in (meter_response, live_response, reading_response, controls_response):
            self.assertContains(response, f'href="{ledger_url}"')
        self.assertContains(meter_response, "Prepaid")
        self.assertContains(live_response, "Top up")
        self.assertContains(live_response, "Refund")
        self.assertContains(detail_response, "Meter Balance Ledger")
        self.assertContains(detail_response, "meter-ledger-table")

    def test_general_settings_redirect_to_combined_parameter_page(self):
        response = self.client.get(reverse("smart_meter:meter_settings"))
        self.assertRedirects(response, reverse("smart_meter:prepaid_params"))
        combined = self.client.get(reverse("smart_meter:prepaid_params"))
        self.assertContains(combined, "General settings")
        self.assertContains(combined, "Prepaid Parameter 1")

    def test_both_ledgers_number_days_and_paginate_without_splitting_a_day(self):
        today = timezone.localdate()
        for offset in range(1, 10):
            day = today - timedelta(days=offset)
            for hour in (10, 11):
                MeterReading.objects.create(
                    meter=self.meter,
                    ts=timezone.make_aware(datetime.combine(day, datetime.min.time()) + timedelta(hours=hour)),
                    balance=Decimal("80.00") + offset,
                    total_energy=Decimal("11.000") + offset,
                    unit_rate=Decimal("50.0000"),
                )
        for url in (
            reverse("smart_meter:prepaid_meter_ledger", args=[self.meter.pk]),
            reverse("smart_meter:meter_detail", args=[self.meter.pk]) + "?tab=ledger",
        ):
            separator = "&" if "?" in url else "?"
            first = self.client.get(url)
            self.assertEqual(first.status_code, 200)
            self.assertEqual(first.context["reading_page"].paginator.num_pages, 2)
            self.assertContains(first, 'data-day-number="1" data-sub-number="1"')
            self.assertContains(first, "ledger-reading-extra d-none")
            second = self.client.get(url + separator + "reading_page=2")
            self.assertEqual(second.status_code, 200)
            self.assertEqual(second.context["reading_page"].number, 2)
            self.assertContains(second, 'data-day-number="8" data-sub-number="2"')
            self.assertContains(second, "Page 2 of 2")

    def test_custom_period_filters_both_ledgers_and_exports(self):
        yesterday = timezone.localdate() - timedelta(days=1)
        MeterReading.objects.create(
            meter=self.meter,
            ts=timezone.make_aware(datetime.combine(yesterday, datetime.min.time()) + timedelta(hours=12)),
            balance=Decimal("88.00"),
            total_energy=Decimal("9.000"),
        )
        query = f"period=custom&from_date={yesterday}&to_date={yesterday}"
        prepaid = self.client.get(reverse("smart_meter:prepaid_meter_ledger", args=[self.meter.pk]) + "?" + query)
        detail = self.client.get(reverse("smart_meter:meter_detail", args=[self.meter.pk]) + "?tab=ledger&" + query)
        self.assertEqual(len(prepaid.context["readings"]), 1)
        self.assertEqual(len(detail.context["ledger_readings"]), 1)
        self.assertContains(detail, 'value="' + str(yesterday) + '"')
        for format, content_type in (
            ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("pdf", "application/pdf"),
            ("jpg", "image/jpeg"),
        ):
            response = self.client.get(reverse("smart_meter:meter_ledger_export", args=[self.meter.pk, format]) + "?" + query)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Content-Type"], content_type)
            self.assertGreater(len(response.content), 100)
            if format == "xlsx":
                workbook = load_workbook(BytesIO(response.content), read_only=True)
                self.assertEqual(workbook["Historical movements"].max_row, 2)
                self.assertEqual(workbook["Historical movements"]["C2"].value, 88)
